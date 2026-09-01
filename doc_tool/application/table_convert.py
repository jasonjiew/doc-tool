# -*- coding: utf-8 -*-
"""表格互转后端：XLSX/CSV 的读取与 →CSV/→XLSX/→Word-HTML 输出。

纯 Python（openpyxl + 标准库 csv），不需要本机 Word；HTML 形态供两条链路
复用：直接落盘成 .html 产物，或经临时 HTML 交给 Word 导出 PDF（A 组的
XLSX/CSV → PDF 走后者）。

既定约束与取舍：

- 打开工作簿必须 ``data_only=True``：公式单元格取缓存值，否则拿到的是公式串。
- 合并单元格无法在 CSV/表格网格中表达：按左上角值填充整个合并区，并在
  ``merged_cells`` 里计数供产物说明提示。
- 损坏/非表格文件折算成结构化 ``TableConvertError``（E6006），绝不抛裸异常。
- CSV 编码探测复用 ``text_convert.read_text_file``（UTF-8 → GBK 兜底）；
  写 CSV 一律 UTF-8 with BOM，Excel 双击打开不乱码。
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from doc_tool.domain.errors import TableConvertError


@dataclass
class TableSheet:
    """单个工作表的展平网格（全部为字符串，空单元格为空串）。"""

    name: str
    rows: List[List[str]] = field(default_factory=list)


@dataclass
class TableBook:
    """一次表格读取的结果：一个或多个工作表 + 合并单元格计数。"""

    sheets: List[TableSheet] = field(default_factory=list)
    merged_cells: int = 0

    @property
    def total_rows(self) -> int:
        return sum(len(sheet.rows) for sheet in self.sheets)

    def summary(self) -> str:
        parts = [
            "「{0}」{1} 行".format(sheet.name, len(sheet.rows))
            for sheet in self.sheets
        ]
        text = "{0} 个工作表：{1}".format(len(self.sheets), "、".join(parts))
        if self.merged_cells:
            text += "；合并单元格 {0} 处已按左上角值填充".format(self.merged_cells)
        return text


def _cell_text(value) -> str:
    """单元格值转显示文本：整数浮点去尾、None 转空串、日期转 ISO。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return str(value).strip()


def _read_xlsx(source: Path) -> TableBook:
    from openpyxl import load_workbook

    book = TableBook()
    try:
        workbook = load_workbook(str(source), data_only=True)
    except Exception as exc:
        raise TableConvertError(details={"source": source.name}) from exc
    try:
        for worksheet in workbook.worksheets:
            merged = 0
            for range_obj in list(worksheet.merged_cells.ranges):
                top_left = worksheet.cell(
                    range_obj.min_row, range_obj.min_col
                ).value
                # 合并区内的从属格是只读 MergedCell：必须先解除合并再回填值。
                worksheet.unmerge_cells(str(range_obj))
                merged += 1
                for row in range(range_obj.min_row, range_obj.max_row + 1):
                    for col in range(range_obj.min_col, range_obj.max_col + 1):
                        worksheet.cell(row, col).value = top_left
            rows = [
                [_cell_text(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            while rows and not any(rows[-1]):
                rows.pop()
            if rows:
                book.sheets.append(TableSheet(name=worksheet.title, rows=rows))
            book.merged_cells += merged
    finally:
        workbook.close()
    if not book.sheets:
        raise TableConvertError(details={"source": source.name})
    return book


def _sniff_delimiter(sample_line: str) -> str:
    """按首行各候选分隔符的出现次数取最可能的（内网 CSV 常见逗号/制表符/分号）。"""
    counts = {symbol: sample_line.count(symbol) for symbol in (",", "\t", ";")}
    best = max(counts, key=lambda symbol: counts[symbol])
    return best if counts[best] > 0 else ","


def _read_csv(source: Path) -> TableBook:
    from doc_tool.application.text_convert import read_text_file

    try:
        text = read_text_file(source)
    except TableConvertError:
        raise
    lines = text.splitlines()
    sample = next((line for line in lines if line.strip()), "")
    if not sample:
        raise TableConvertError(details={"source": source.name})
    reader = csv.reader(lines, delimiter=_sniff_delimiter(sample))
    rows = [[_cell_text(cell) for cell in row] for row in reader if row]
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        raise TableConvertError(details={"source": source.name})
    return TableBook(sheets=[TableSheet(name=source.stem, rows=rows)])


def read_table_book(source: Path) -> TableBook:
    """按扩展名读取表格文件；不支持的扩展名与损坏文件都折算成 E6006。"""
    source = Path(source)
    suffix = source.suffix.lower()
    if suffix == ".xlsx":
        return _read_xlsx(source)
    if suffix == ".csv":
        return _read_csv(source)
    raise TableConvertError(details={"source": source.name})


def sheet_to_csv(sheet: TableSheet, target: Path) -> None:
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(str(target), "w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(sheet.rows)


def rows_to_xlsx(rows: List[List[str]], target: Path, sheet_name: str = "Sheet1") -> None:
    from openpyxl import Workbook

    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31] or "Sheet1"
    for row in rows:
        worksheet.append(row)
    workbook.save(str(target))


def _escape(text: str) -> str:
    import html

    return html.escape(text, quote=False)


def book_to_word_html(book: TableBook, title: str = "") -> str:
    """把工作簿渲染成完整 HTML（内嵌与 Markdown→Word 同一套排版样式）。

    产物有两个去向：直接落盘为 .html，或写临时文件交 Word 导出 PDF——
    样式沿用 ``markdown_word.WORD_STYLE_CSS``，两条链路的表格观感一致。
    每个工作表一个一级标题 + 表格；首行按表头渲染。
    """
    import html

    from doc_tool.application.markdown_word import WORD_STYLE_CSS

    blocks: List[str] = []
    for sheet in book.sheets:
        blocks.append("<h1>{0}</h1>".format(_escape(sheet.name)))
        if not sheet.rows:
            blocks.append("<p>（空工作表）</p>")
            continue
        header, body = sheet.rows[0], sheet.rows[1:]
        parts = ["<table>", "<thead><tr>"]
        parts.extend("<th>{0}</th>".format(_escape(cell)) for cell in header)
        parts.append("</tr></thead>")
        if body:
            parts.append("<tbody>")
            for row in body:
                parts.append("<tr>")
                parts.extend("<td>{0}</td>".format(_escape(cell)) for cell in row)
                parts.append("</tr>")
            parts.append("</tbody>")
        parts.append("</table>")
        blocks.append("".join(parts))
    head = "<title>{0}</title>".format(html.escape(title, quote=True)) if title else ""
    return (
        '<html xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:w="urn:schemas-microsoft-com:office:word" '
        'xmlns="http://www.w3.org/TR/REC-html40">'
        '<head>{head}<meta charset="utf-8"/>'
        "<style>{css}</style></head><body>{body}</body></html>"
    ).format(head=head, css=WORD_STYLE_CSS, body="".join(blocks))


def write_book_csv(book: TableBook, target: Path) -> str:
    """XLSX→CSV：只导第一个工作表（CSV 没有多表概念），返回产物说明。"""
    sheet = book.sheets[0]
    sheet_to_csv(sheet, target)
    detail = "工作表「{0}」{1} 行已导出".format(sheet.name, len(sheet.rows))
    if len(book.sheets) > 1:
        detail += "；其余 {0} 个工作表未导出（CSV 为单表格式）".format(len(book.sheets) - 1)
    return detail


def write_rows_xlsx(rows: List[List[str]], target: Path, sheet_name: str) -> str:
    rows_to_xlsx(rows, target, sheet_name)
    return "{0} 行 × {1} 列已写入工作表「{2}」".format(
        len(rows), max(len(row) for row in rows), sheet_name
    )
