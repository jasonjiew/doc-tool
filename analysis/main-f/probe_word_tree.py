# -*- coding: utf-8 -*-
from __future__ import annotations
import shutil, sys, tempfile
from pathlib import Path
REPO = Path(__file__).resolve().parents[2]
for c in (str(REPO), str(REPO / "scripts")):
    if c not in sys.path:
        sys.path.insert(0, c)
from docx import Document
from doc_tool.application.intake_entries import run_intake
from doc_tool.application.effective_snapshot import discover_chapters
from doc_tool.domain.manifest import ProjectManifest
from doc_tool.domain.paths import ProjectPaths
tmp = Path(tempfile.mkdtemp(prefix="main-f-probe-"))
img = tmp / "fig.png"
from PIL import Image
Image.new("RGB", (20, 14), (1, 2, 3)).save(str(img))
src = tmp / "doc.docx"
d = Document()
d.add_heading("引言", level=1)
d.add_paragraph("第一轮正文。")
d.add_heading("目的", level=2)
d.add_paragraph("目的正文。")
d.add_picture(str(img))
d.add_heading("设计", level=1)
d.add_paragraph("设计正文。")
d.save(str(src))
out = run_intake(src, parent_dir=tmp, target_name="proj")
print("ok", out.ok, out.errors)
root = Path(out.project_root)
for p in sorted(root.rglob("*")):
    if p.is_file():
        print("  ", p.relative_to(root).as_posix())
man = ProjectManifest.load(root)
paths = ProjectPaths(root)
cr = paths.resolve(man.relative_content_root())
print("discover:", [r for r, _ in discover_chapters(cr)])
print("manifest.chapters:", man.chapters)
shutil.rmtree(tmp, ignore_errors=True)