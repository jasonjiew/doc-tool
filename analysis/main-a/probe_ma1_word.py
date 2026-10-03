# -*- coding: utf-8 -*-
"""MAIN-A 1.1 probe: mixed Word import read-back."""
from __future__ import annotations

import hashlib
import shutil
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "scripts"))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from lxml import etree  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _png(path: Path) -> Path:
    from PIL import Image
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (24, 16), (10, 120, 200)).save(str(path))
    return path


def _append_raw(document, xml: str) -> None:
    body = document.element.body
    element = etree.fromstring(
        '<w:p xmlns:w="{0}" xmlns:m="{1}" xmlns:r="{2}">{3}</w:p>'.format(W_NS, M_NS, R_NS, xml)
    )
    body.insert(len(body) - 1, element)


def build_mixed(path: Path, image: Path) -> Path:
    document = Document()
    document.styles["Normal"].font.size = Pt(10.5)
    document.add_heading("引言", level=1)
    document.add_paragraph("本文档说明导入主流程。")
    document.add_heading("目的", level=2)
    document.add_paragraph("验证列表、表格、图片与代码的读回。")

    # 项目符号列表（真实 Word 列表段落带直接 w:numPr）
    for text in ("第一项", "第二项"):
        document.add_paragraph(text)
        para = document.paragraphs[-1]
        pPr = para._p.get_or_add_pPr()
        pPr.append(etree.fromstring(
            '<w:numPr xmlns:w="{0}"><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'.format(W_NS)
        ))
    # 编号列表：文本自带编号（真实 Word 常见）
    document.add_paragraph("1. 步骤一")
    document.add_paragraph("2. 步骤二")

    document.add_heading("表格与图", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "前导零"
    table.cell(1, 1).text = "007"
    document.add_paragraph("表格后正文。")
    document.add_picture(str(image))
    document.add_paragraph("图后正文。")

    document.add_heading("代码与链接", level=2)
    code = document.add_paragraph("def f():")
    code.style = document.styles["Normal"]
    document.add_paragraph("    return 1")
    _append_raw(
        document,
        '<w:hyperlink r:id="rId99"><w:r><w:t>外链文字</w:t></w:r></w:hyperlink>',
    )
    document.add_paragraph("行内 address = 0x00 保持原样。")
    _append_raw(document, '<w:r><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath></w:r>')
    _append_raw(
        document,
        '<w:r><w:txbxContent><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p></w:txbxContent></w:r>',
    )
    document.add_paragraph("结束段。")
    document.save(str(path))
    return path


tmp = Path(tempfile.mkdtemp(prefix="main-a-word-"))
image = _png(tmp / "fig.png")
source = build_mixed(tmp / "混合样例.docx", image)
before = hashlib.sha256(source.read_bytes()).hexdigest()

from doc_tool.application.intake_entries import run_intake  # noqa: E402

outcome = run_intake(source, parent_dir=tmp, target_name="混合项目")
print("ok=", outcome.ok)
for err in outcome.errors:
    print("   err:", err)
for warn in outcome.warnings[:10]:
    print("   warn:", warn)
root = Path(outcome.project_root) if outcome.project_root else None
if root is not None:
    print("--- content tree ---")
    for path in sorted(root.rglob("*")):
        if path.is_file():
            rel = path.relative_to(root).as_posix()
            if rel.startswith(("content/", "assets/")):
                print("===", rel)
                if rel.endswith(".md"):
                    print(path.read_text(encoding="utf-8"))
    print("source unchanged=", hashlib.sha256(source.read_bytes()).hexdigest() == before)
    print("original retained=", (root / "original" / "source.docx").is_file())
print("tmp=", tmp)