# -*- coding: utf-8 -*-
"""MAIN-A 1.2 probe: preservation findings for mixed Word import."""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(REPO / "scripts" / "tests"))

from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _append_raw(document, xml):
    body = document.element.body
    element = etree.fromstring('<w:p xmlns:w="{0}" xmlns:m="{1}">{2}</w:p>'.format(W_NS, M_NS, xml))
    body.insert(len(body) - 1, element)


tmp = Path(tempfile.mkdtemp(prefix="main-a-find-"))
source = tmp / "复杂.docx"
document = Document()
document.add_heading("第一章", level=1)
document.add_paragraph("正文。")
_append_raw(document, '<w:r><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath></w:r>')
_append_raw(document, '<w:r><w:txbxContent><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p></w:txbxContent></w:r>')
_append_raw(document, '<w:ins w:id="1" w:author="a" w:date="2026-10-01T00:00:00Z"><w:r><w:t>插入的修订文字</w:t></w:r></w:ins>')
document.add_paragraph("结束。")
document.save(str(source))

from doc_tool.application.intake_entries import run_intake  # noqa: E402

outcome = run_intake(source, parent_dir=tmp, target_name="复杂项目")
print("ok=", outcome.ok, "errors=", outcome.errors)
plan = outcome.plan
print("--- heading decisions ---")
for item in plan.headingDecisions:
    print("   ", item.title, item.level, item.source, item.action, item.source_index)
print("--- preservation findings ---")
for item in plan.preservationFindings:
    print("   ", item.feature, "|", item.handling, "|retained=", item.retained_path, "|chapter=", item.target_chapter, "|path=", item.target_path, "|line=", item.target_line, "|action=", item.action)
print("--- unavailable comparisons ---")
print("   ", plan.unavailableComparisons)
root = Path(outcome.project_root)
md = list((root / "content").rglob("*.md"))
for path in md:
    if path.name != "_revision_record.md":
        print("===", path.relative_to(root).as_posix())
        print(path.read_text(encoding="utf-8"))
shutil.rmtree(tmp, ignore_errors=True)