# -*- coding: utf-8 -*-
"""Phase 2: 建立 Word 模板骨架（薄壳，委托给 doc_tool.adapters.importer）。

用法:
    python scripts/migration/make_template.py requirement
    python scripts/migration/make_template.py design
    python scripts/migration/make_template.py all

策略: 复制原始 Word -> 删除正文可变内容(从第一个 Heading1 起) -> 保留封面/TOC/
修订记录/页眉页脚/Section/样式。原始 Word 永不被修改。

任务 4.1：实际逻辑已重构为 ``generate_template`` 无全局副作用纯函数，本脚本
仅作为旧 CLI 入口的兼容薄壳。
"""
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC_DIR = os.path.dirname(BASE)

DOCS = {
    "requirement": {
        "src": "KF-2090-1-001 康尚健康云软件需求说明书(3.8).docx",
        "out": os.path.join(BASE, "templates", "requirement-template.docx"),
    },
    "design": {
        "src": "KF-2090-1-006 康尚健康云系统详细设计说明书(2.5).docx",
        "out": os.path.join(BASE, "templates", "design-template.docx"),
    },
}

sys.path.insert(0, BASE)
from doc_tool.adapters.importer import generate_template  # noqa: E402


def build_template(key):
    cfg = DOCS[key]
    src = os.path.join(SRC_DIR, cfg["src"])
    out = cfg["out"]
    if not os.path.exists(src):
        print("MISSING:", src)
        return False
    print("[%s] 源文档: %s" % (key, src))
    meta = generate_template(src, out)
    print("[%s] 正文起点: body 子元素索引 %d" % (key, meta.body_start_index))
    print("[%s] 模板已生成: %s (%.2f MB)" % (key, out, os.path.getsize(out) / 1024 / 1024))
    return True


def verify(key):
    """用 python-docx 尝试打开，确认无结构损坏"""
    import docx
    from lxml import etree
    import zipfile

    cfg = DOCS[key]
    out = cfg["out"]
    d = docx.Document(out)
    print("[%s] 校验通过: python-docx 可打开, 段落=%d 表格=%d" % (key, len(d.paragraphs), len(d.tables)))
    with zipfile.ZipFile(out) as z:
        etree.fromstring(z.read("word/document.xml"))
    print("[%s] XML 良构校验通过" % key)


if __name__ == "__main__":
    arg = sys.argv[1] if len(sys.argv) > 1 else "all"
    keys = list(DOCS) if arg == "all" else [arg]
    for k in keys:
        if build_template(k):
            verify(k)
