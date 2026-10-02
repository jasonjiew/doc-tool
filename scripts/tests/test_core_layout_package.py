# -*- coding: utf-8 -*-
"""CORE-G 排版与源码包测试（7.1-7.5 的自动部分）。"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.tests import core_fixtures as fixtures  # noqa: E402


def _build_docx_with_media(path: Path, image: Path) -> Path:
    """构造含大图 + 普通表格 + 两章的 DOCX（不依赖 Word）。"""
    from docx import Document
    from docx.shared import Mm

    document = Document()
    document.add_heading("第1章 总述", level=1)
    document.add_paragraph("总述正文。")
    document.add_picture(str(image), width=Mm(400))  # 明显超宽
    table = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = "R{0}C{1}".format(row_index, col_index)
    document.add_heading("第2章 详述", level=1)
    document.add_paragraph("详述正文。")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.work = fixtures.scratch_dir("layout")
        self.image = fixtures.tiny_png(self.work / "img.png")
        self.docx = _build_docx_with_media(self.work / "doc.docx", self.image)

    def tearDown(self):
        fixtures.cleanup(self.work)

    def test_body_adaptive_scales_images_and_formats_tables(self):
        from docx import Document
        from doc_tool.application.export.layout_profile import (
            LayoutProfile, apply_layout_to_docx,
        )

        profile = LayoutProfile.body_adaptive(page_break_before_chapter=True)
        outcome = apply_layout_to_docx(self.docx, profile, chapter_titles=["第1章 总述", "第2章 详述"])
        self.assertTrue(outcome.ok, outcome.message)
        self.assertEqual(outcome.images_scaled, 1)
        self.assertEqual(outcome.tables_formatted, 1)
        self.assertEqual(outcome.headers_repeated, 1)
        self.assertGreaterEqual(outcome.page_breaks, 1)
        document = Document(str(self.docx))
        section = document.sections[0]
        limit = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
        for shape in document.inline_shapes:
            self.assertLessEqual(int(shape.width), limit)
        # 重复表头已写入 OOXML
        from docx.oxml.ns import qn

        header = document.tables[0].rows[0]._tr.find(qn("w:trPr"))
        self.assertIsNotNone(header)
        self.assertIsNotNone(header.find(qn("w:tblHeader")))

    def test_template_mode_does_not_modify_document(self):
        from doc_tool.application.export.layout_profile import (
            LayoutProfile, apply_layout_to_docx,
        )

        before = self.docx.read_bytes()
        outcome = apply_layout_to_docx(self.docx, LayoutProfile())
        self.assertTrue(outcome.ok)
        self.assertEqual(self.docx.read_bytes(), before)
        self.assertIn("沿用底模", outcome.summary_lines()[0])

    def test_complex_table_is_kept_as_is(self):
        from docx import Document
        from doc_tool.application.export.layout_profile import (
            LayoutProfile, apply_layout_to_docx,
        )

        document = Document(str(self.docx))
        table = document.tables[0]
        table.cell(0, 0).merge(table.cell(1, 0))  # 制造合并单元格
        document.save(str(self.docx))
        outcome = apply_layout_to_docx(
            self.docx, LayoutProfile.body_adaptive(), chapter_titles=["第1章 总述"],
        )
        self.assertTrue(outcome.ok)
        self.assertEqual(outcome.tables_skipped_complex, 1)
        self.assertEqual(outcome.tables_formatted, 0)
        self.assertTrue(any("复杂表格" in line for line in outcome.summary_lines()))

    def test_landscape_without_own_section_is_reported_not_forced(self):
        from doc_tool.application.export.layout_profile import (
            LayoutProfile, apply_layout_to_docx,
        )

        profile = LayoutProfile(
            mode="body-adaptive", landscape_chapters=["第2章 详述"],
        )
        outcome = apply_layout_to_docx(self.docx, profile, chapter_titles=["第1章 总述", "第2章 详述"])
        self.assertTrue(outcome.ok)
        self.assertTrue(outcome.landscape_skipped or outcome.landscape_applied)
        for line in outcome.summary_lines():
            self.assertNotIn("隐藏", line)


class ExportLayoutIntegrationTests(unittest.TestCase):
    """出稿链路：body-adaptive 通过 ExportRequest.layout 生效。"""

    _shared_root = None
    _shared_project = None

    @classmethod
    def setUpClass(cls):
        cls._shared_root = fixtures.scratch_dir("layout-export")
        cls._shared_project = fixtures.two_chapter_project(cls._shared_root / "proj")

    @classmethod
    def tearDownClass(cls):
        fixtures.cleanup(cls._shared_root)

    def setUp(self):
        import shutil

        self.work = fixtures.scratch_dir("layout-out")
        self.project = self.work / "proj"
        shutil.copytree(str(self._shared_project), str(self.project))
        self.out = self.work / "out"
        self.out.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        fixtures.cleanup(self.work)

    def test_export_applies_body_adaptive_layout(self):
        from doc_tool.application.intake_contract import (
            FORMAT_DOCX, ExportRequest, SOURCE_MODE_SAVED,
        )
        from doc_tool.application.project_export import run_project_export

        request = ExportRequest(
            project_root=str(self.project), formats=[FORMAT_DOCX],
            source_mode=SOURCE_MODE_SAVED, destination=str(self.out),
            layout_profile="body-adaptive",
        )
        report = run_project_export(request, skip_word_refresh=True)
        docx = report.result_for(FORMAT_DOCX)
        self.assertTrue(docx.usable, docx.message)
        self.assertTrue(Path(docx.path).is_file())
        self.assertTrue(any("排版" in item or "表格" in item or "图片" in item
                            for item in docx.warnings), docx.warnings)


if __name__ == "__main__":
    unittest.main(verbosity=2)