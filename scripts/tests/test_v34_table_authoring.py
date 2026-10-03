# -*- coding: utf-8 -*-
"""V3.4 普通表格编辑回归（34-A～34-E）。

覆盖表格模型往返、非矩形补齐、多行表达、片段身份、网格应用（一次撤销、
源变化兜底）、TSV/管道表格粘贴、预览与 CORE 同源出稿、大输入测量与键盘可达。
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[2]
for candidate in (str(ROOT), str(ROOT / "scripts"), str(ROOT / "scripts" / "tests")):
    if candidate not in sys.path:
        sys.path.insert(0, candidate)

from scripts.tests import core_fixtures as fixtures  # noqa: E402
from doc_tool.application.content import table_grid as grid  # noqa: E402
from doc_tool.domain.manifest import ProjectManifest  # noqa: E402

EVIDENCE_DIR = ROOT / "analysis" / "product-v34-table-20261003"


def _cleanup(path):
    assert Path(path).resolve().is_relative_to(fixtures.SCRATCH_ROOT.resolve())
    fixtures.cleanup(path)


def _write_evidence(name: str, payload: dict) -> None:
    if os.environ.get("PRODUCT_V34_EVIDENCE") != "1":
        return
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    (EVIDENCE_DIR / name).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8",
    )


SAMPLE_TABLE = (
    "| 参数 | 说明 | 空列 |\n"
    "| :--- | ---: | --- |\n"
    "| a\\|b | 反斜杠 \\\\ 与中文 |  |\n"
    "| 0012 | =SUM(A1:A2) | 2026-10-01 |\n"
)


class TableModelTests(unittest.TestCase):
    """34-A：模型、表达与来源身份。"""

    def test_roundtrip_keeps_pipes_backslash_empty_cells_and_alignment(self):
        parsed = grid.parse_table(SAMPLE_TABLE)
        self.assertTrue(parsed.ok, parsed.error)
        model = parsed.model
        self.assertEqual(model.header, ["参数", "说明", "空列"])
        self.assertEqual(model.row_count, 2)
        self.assertEqual(model.alignments[:2], ["left", "right"])
        self.assertEqual(model.rows[0][0], "a|b")
        self.assertEqual(model.rows[0][1], "反斜杠 \\ 与中文")
        self.assertEqual(model.rows[0][2], "", "空单元格必须保留")
        self.assertEqual(model.rows[1][0], "0012", "前导零必须保留")
        self.assertEqual(model.rows[1][1], "=SUM(A1:A2)", "公式文本不得求值")
        again = grid.parse_table(model.serialize())
        self.assertTrue(again.ok, again.error)
        self.assertEqual(again.model.header, model.header)
        self.assertEqual(again.model.rows, model.rows)
        self.assertEqual(again.model.alignments, model.alignments)

    def test_trailing_empty_columns_are_not_dropped(self):
        model = grid.parse_table("| a | b | c |\n| --- | --- | --- |\n| 1 |  |  |\n").model
        self.assertEqual(model.column_count, 3)
        self.assertEqual(len(model.rows[0]), 3)
        self.assertEqual(model.rows[0][1:], ["", ""])

    def test_ragged_input_is_padded_with_warning_and_raw_preserved(self):
        parsed = grid.parse_table("| a | b | c |\n| --- | --- | --- |\n| 1 | 2 |\n")
        self.assertTrue(parsed.ok)
        model = parsed.model
        self.assertTrue(model.warnings, "非矩形输入必须给出补齐提示")
        self.assertEqual(len(model.rows[0]), 3)
        self.assertEqual(model.rows[0][2], "")
        self.assertIn("| 1 | 2 |", model.raw_input)

    def test_non_table_and_broken_separator_are_rejected_with_reason(self):
        self.assertFalse(grid.parse_table("普通段落\n没有表格\n").ok)
        parsed = grid.parse_table("| a | b |\n| 不是分隔行 | x |\n| 1 | 2 |\n")
        self.assertFalse(parsed.ok)
        self.assertIn("分隔行", parsed.error)

    def test_multiline_cells_use_supported_marker_in_preview_and_word(self):
        preview_html = None
        docx_cell = None
        work = fixtures.scratch_dir("v34-multiline")
        try:
            from doc_tool.application.content.preview import render_markdown_html

            table = "| 步骤 | 说明 |\n| --- | --- |\n| 1 | 第一行{0}第二行 |\n".format(
                grid.MULTILINE_MARKER,
            )
            preview_html = render_markdown_html(table)
            self.assertIn("<br/>", preview_html, "预览应保留 <br/> 换行表达")

            from doc_tool.application.intake_contract import (
                FORMAT_DOCX, SOURCE_MODE_SAVED, ExportRequest,
            )
            from doc_tool.application.project_export import run_project_export

            # 用带有效底模的两章项目夹具，保证诊断构建能真正产出 DOCX。
            project = fixtures.two_chapter_project(work / "表格工程")
            manifest = ProjectManifest.load(project)
            content = project / manifest.relative_content_root()
            chapters = [
                path for path in sorted(content.rglob("*.md")) if not path.name.startswith("_")
            ]
            self.assertTrue(chapters, "夹具项目应包含章节")
            chapters[0].write_text(
                chapters[0].read_text(encoding="utf-8") + "\n" + table, encoding="utf-8",
            )
            report = run_project_export(
                ExportRequest(
                    project_root=str(project), formats=[FORMAT_DOCX],
                    source_mode=SOURCE_MODE_SAVED, destination=str(work / "out"),
                ),
                skip_word_refresh=True,
            )
            docx_result = report.result_for(FORMAT_DOCX)
            self.assertTrue(docx_result is not None and docx_result.usable, "诊断构建应产出 DOCX")
            from docx import Document

            document = Document(docx_result.path)
            self.assertTrue(document.tables, "出稿应包含表格")
            cell_texts = [
                cell.text for table_obj in document.tables for row in table_obj.rows for cell in row.cells
            ]
            docx_cell = [text for text in cell_texts if "第一行" in text]
            self.assertTrue(docx_cell, "Word 表格应保留多行单元格内容")
            self.assertIn("第二行", docx_cell[0].replace("\n", ""))
        finally:
            _cleanup(work)
        _write_evidence("multiline-support.json", {
            "marker": grid.MULTILINE_MARKER,
            "preview": "ok" if preview_html and "<br/>" in preview_html else "check",
            "wordCell": docx_cell[0] if docx_cell else "",
            "platform": "Qt offscreen + 诊断 DOCX 构建",
        })

    def test_fragment_identity_unique_change_and_ambiguity(self):
        source = "前言\n\n" + SAMPLE_TABLE + "\n后记\n"
        located = grid.locate_fragment(source, SAMPLE_TABLE)
        self.assertTrue(located.ok)
        updated, reason = grid.apply_fragment(source, SAMPLE_TABLE, "| 参数 |\n| --- |\n| 新值 |")
        self.assertEqual(reason, "")
        self.assertIn("新值", updated)
        self.assertIn("后记", updated)
        changed = source.replace("| 0012 |", "| 9999 |")
        missing, reason = grid.apply_fragment(changed, SAMPLE_TABLE, "x")
        self.assertIsNone(missing)
        self.assertIn("已不在当前缓冲", reason)
        ambiguous_source = SAMPLE_TABLE + "\n\n" + SAMPLE_TABLE + "\n"
        blocked, reason = grid.apply_fragment(ambiguous_source, SAMPLE_TABLE, "x")
        self.assertIsNone(blocked)
        self.assertIn("无法唯一确定", reason)
        resolved, _reason = grid.apply_fragment(ambiguous_source, SAMPLE_TABLE, "x", near_line=5)
        self.assertIsNone(resolved, "距离不代表片段身份；应由用户明确选择目标表格")


class DelimitedIntakeTests(unittest.TestCase):
    """34-C：TSV/管道表格解析与粘贴预览。"""

    def test_tsv_quotes_leading_zeros_dates_and_formulas(self):
        text = '名称\t数量\t日期\t公式\n"含\t制表""符"\t0012\t2026-10-01\t=SUM(A1:A2)\n'
        result = grid.parse_clipboard_text(text)
        self.assertTrue(result.ok, result.error)
        self.assertEqual(result.source, "tsv")
        self.assertEqual(result.rows[1][1], "0012")
        self.assertEqual(result.rows[1][2], "2026-10-01")
        self.assertEqual(result.rows[1][3], "=SUM(A1:A2)")
        self.assertIn("含\t制表\"符", result.rows[1][0])
        serialized = grid.model_from_matrix(result.rows).serialize()
        self.assertIn("0012", serialized)
        self.assertIn("=SUM(A1:A2)", serialized)

    def test_tsv_quoted_newline_uses_supported_marker(self):
        result = grid.parse_clipboard_text('说明\t值\n"第一行\n第二行"\t1\n')
        self.assertTrue(result.ok, result.error)
        self.assertEqual(result.multi_line_cells, 1)
        self.assertIn(grid.MULTILINE_MARKER, result.rows[1][0])
        self.assertTrue(any("换行" in item for item in result.warnings))

    def test_tsv_ragged_rows_are_padded_and_reported(self):
        result = grid.parse_clipboard_text("a\tb\tc\n1\t2\n")
        self.assertTrue(result.ok)
        self.assertEqual(result.column_count, 3)
        self.assertEqual(result.rows[1][2], "")
        self.assertTrue(result.warnings)

    def test_markdown_input_and_unsupported_input(self):
        parsed = grid.parse_clipboard_text(SAMPLE_TABLE)
        self.assertTrue(parsed.ok)
        self.assertEqual(parsed.source, "markdown")
        unsupported = grid.parse_clipboard_text("就是一句普通文本\n没有制表符也没有竖线\n")
        self.assertFalse(unsupported.ok)
        self.assertIn("普通粘贴", unsupported.error)

    def test_preview_pages_large_input_without_truncating_insert(self):
        rows = [["r{0}c{1}".format(r, c) for c in range(5)] for r in range(100)]
        text = "\n".join("\t".join(row) for row in rows)
        result = grid.parse_clipboard_text(text)
        self.assertTrue(result.ok)
        shown, truncated, total = grid.preview_rows(result.rows)
        self.assertEqual(total, 100)
        self.assertTrue(truncated)
        self.assertEqual(len(shown), grid.PREVIEW_ROW_LIMIT)
        serialized = grid.model_from_matrix(result.rows).serialize()
        self.assertIn("r99c4", serialized, "应用必须包含完整输入")


class TableGridGuiTests(unittest.TestCase):
    """34-B/34-C：网格对话框与编辑器接入（真实 Qt 控件）。"""

    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from doc_tool.ui.content.editor_panel import EditorPanel

        self._writer_writes = []

        class _Writer:
            def write(self, rel_path, text):
                self._writer_writes.append((rel_path, text))

            def write_text(self, rel_path, text):
                self._writer_writes.append((rel_path, text))
                return True

            def resolve(self, rel_path):
                return Path.cwd() / str(rel_path)

            def __getattr__(self, name):
                raise AttributeError(name)

        self.source_text = "前言\n\n" + SAMPLE_TABLE + "\n后记\n"
        self.panel = EditorPanel(_Writer())
        self.panel.load("requirement/1 概述/1.1 背景.md", self.source_text)
        self._put_cursor_on_table()

    def _put_cursor_on_table(self) -> None:
        """把光标放到表格第一行：真实用户是“光标在表格内”才打开网格。"""
        document = self.panel._editor.document()
        block = document.findBlockByNumber(2)
        cursor = self.panel._editor.textCursor()
        cursor.setPosition(block.position())
        self.panel._editor.setTextCursor(cursor)

    def tearDown(self):
        self.panel.deleteLater()
        self.app.processEvents()

    def _model(self):
        parsed = grid.parse_table(SAMPLE_TABLE)
        self.assertTrue(parsed.ok)
        return parsed.model

    def test_grid_edits_rows_columns_and_alignment(self):
        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        dialog = TableGridDialog(self._model())
        try:
            dialog.resize(1280, 720)
            dialog.show()
            self.app.processEvents()
            self.assertEqual(dialog.table.rowCount(), 3, "表头 + 2 数据行")
            self.assertEqual(dialog.table.columnCount(), 3)
            dialog.table.setCurrentCell(1, 0)
            dialog.add_row_btn.click()
            self.assertEqual(dialog.table.rowCount(), 4)
            dialog.table.setCurrentCell(2, 0)
            dialog.remove_row_btn.click()
            self.assertEqual(dialog.table.rowCount(), 3)
            dialog.table.setCurrentCell(1, 1)
            dialog.move_col_left_btn.click()
            dialog.table.setCurrentCell(1, 1)
            dialog.align_buttons["center"].click()
            self.assertEqual(dialog.result_model().alignments[1], "center")
            text = dialog.markdown_text()
            self.assertTrue(text.startswith("|"))
            header_line = text.splitlines()[0]
            self.assertIn("参数", header_line)
            self.assertIn("|", text.splitlines()[1])
            self.assertTrue(dialog.is_dirty())
        finally:
            dialog.close()
            dialog.deleteLater()

    def test_grid_local_undo_redo_for_rows_columns_and_cells(self):
        from PySide6.QtWidgets import QTableWidgetItem

        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        dialog = TableGridDialog(self._model())
        try:
            original = dialog.markdown_text()
            dialog.table.setCurrentCell(1, 0)
            dialog.add_row_btn.click()
            self.assertEqual(dialog.table.rowCount(), 4)
            dialog.table.setItem(1, 0, QTableWidgetItem("单元格改动"))
            self.assertIn("单元格改动", dialog.markdown_text())
            dialog.undo_grid()
            self.assertNotIn("单元格改动", dialog.markdown_text())
            dialog.undo_grid()
            self.assertEqual(dialog.table.rowCount(), 3)
            self.assertEqual(dialog.markdown_text(), original)
            dialog.redo_grid()
            self.assertEqual(dialog.table.rowCount(), 4)
            dialog.redo_grid()
            self.assertIn("单元格改动", dialog.markdown_text())
        finally:
            dialog.close()
            dialog.deleteLater()

    def test_grid_tab_navigation_and_chinese_cell(self):
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QKeyEvent
        from PySide6.QtWidgets import QApplication

        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        dialog = TableGridDialog(self._model())
        try:
            dialog.show()
            self.app.processEvents()
            dialog.table.setCurrentCell(0, 0)
            for _ in range(3):
                QApplication.sendEvent(
                    dialog.table,
                    QKeyEvent(QKeyEvent.Type.KeyPress, Qt.Key.Key_Tab, Qt.KeyboardModifier.NoModifier),
                )
            self.assertEqual(
                (dialog.table.currentRow(), dialog.table.currentColumn()), (1, 0),
                "Tab 应移动到下一个单元格并跨行",
            )
            from PySide6.QtWidgets import QTableWidgetItem

            dialog.table.setItem(1, 0, QTableWidgetItem("中文新值"))
            self.assertIn("中文新值", dialog.markdown_text())
        finally:
            dialog.close()
            dialog.deleteLater()

    def test_apply_grid_is_single_undo_and_keeps_other_text(self):
        from PySide6.QtWidgets import QDialog, QTableWidgetItem

        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        def fake_exec(dialog_self):
            dialog_self.table.setItem(1, 0, QTableWidgetItem("改名后的参数"))
            return QDialog.DialogCode.Accepted

        with patch.object(TableGridDialog, "exec", fake_exec):
            applied = self.panel.open_table_grid()
        self.assertTrue(applied)
        text = self.panel._editor.toPlainText()
        self.assertIn("改名后的参数", text)
        self.assertIn("前言", text)
        self.assertIn("后记", text, "表格以外的文字必须保留")
        self.panel._editor.undo()
        self.assertEqual(self.panel._editor.toPlainText(), self.source_text)
        self.panel._editor.redo()
        self.assertIn("改名后的参数", self.panel._editor.toPlainText(), "编辑器重做应恢复表格修改")
        self.panel._editor.undo()

    def test_source_changed_keeps_pending_text_without_overwriting(self):
        from PySide6.QtWidgets import QDialog, QTableWidgetItem

        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        calls = {"count": 0}

        def fake_exec(dialog_self):
            calls["count"] += 1
            if calls["count"] == 1:
                dialog_self.table.setItem(1, 0, QTableWidgetItem("网格新值"))
                # 模拟网格打开期间原表格被其他编辑改动
                self.panel._editor.setPlainText("前言\n\n| 参数 | 说明 | 空列 |\n| :--- | ---: | --- |\n| 被外部改过 | x |  |\n后记\n")
                return QDialog.DialogCode.Accepted
            return QDialog.DialogCode.Rejected

        with patch.object(TableGridDialog, "exec", fake_exec):
            applied = self.panel.open_table_grid()
        self.assertFalse(applied)
        current = self.panel._editor.toPlainText()
        self.assertIn("被外部改过", current, "旧草稿不得覆盖已变化的原表格")
        self.assertNotIn("网格新值", current)
        self.assertIn("网格新值", self.panel._pending_table_text, "待应用文本必须保留")

    def test_non_table_line_reports_reason_without_dialog(self):
        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        self.panel._editor.setPlainText("普通段落\n没有表格\n")
        with patch.object(TableGridDialog, "exec", side_effect=AssertionError("不应打开网格")):
            applied = self.panel.open_table_grid()
        self.assertFalse(applied)
        self.assertIn("表格", self.panel._status_label.text())

    def test_paste_as_table_inserts_full_table_and_keeps_plain_paste(self):
        from PySide6.QtWidgets import QDialog

        from doc_tool.ui.content.table_grid_dialog import TablePasteDialog

        rows = "\n".join("\t".join("r{0}c{1}".format(r, c) for c in range(3)) for r in range(30))

        def fake_exec(dialog_self):
            dialog_self.mode = TablePasteDialog.MODE_TABLE
            return QDialog.DialogCode.Accepted

        with patch.object(TablePasteDialog, "exec", fake_exec):
            inserted = self.panel.paste_as_table(rows)
        self.assertTrue(inserted)
        text = self.panel._editor.toPlainText()
        self.assertIn("r29c2", text, "大输入必须完整插入")
        self.assertIn("r0c0", text)
        self.panel._editor.undo()
        self.assertNotIn("r29c2", self.panel._editor.toPlainText())

    def test_paste_cancel_and_raw_fallback_keep_content(self):
        from PySide6.QtWidgets import QDialog

        from doc_tool.ui.content.table_grid_dialog import TablePasteDialog

        original = self.panel._editor.toPlainText()
        with patch.object(TablePasteDialog, "exec", return_value=QDialog.DialogCode.Rejected):
            inserted = self.panel.paste_as_table("a\tb\n1\t2\n")
        self.assertFalse(inserted)
        self.assertEqual(self.panel._editor.toPlainText(), original, "取消不得修改缓冲")

        unsupported = "一句没有表格结构的文本"
        with patch.object(TablePasteDialog, "exec", return_value=QDialog.DialogCode.Rejected):
            inserted = self.panel.paste_as_table(unsupported)
        self.assertFalse(inserted)
        self.assertEqual(self.panel._editor.toPlainText(), original)

    def test_preview_dialog_pages_large_input_and_reports_support(self):
        from doc_tool.ui.content.table_grid_dialog import TablePasteDialog

        rows = "\n".join("\t".join("c{0}".format(c) for c in range(4)) for _ in range(60))
        result = grid.parse_clipboard_text(rows)
        dialog = TablePasteDialog(result)
        try:
            dialog.resize(1280, 720)
            dialog.show()
            self.app.processEvents()
            self.assertEqual(dialog.table.rowCount(), grid.PREVIEW_ROW_LIMIT)
            self.assertIn("共 60 行", dialog.page_label.text())
            self.assertIn("不截断", dialog.page_label.text())
            self.assertIn("完整插入", dialog.page_label.text())
            self.assertIn("c3", dialog.markdown_text())
            dialog.header_check.setChecked(False)
            self.app.processEvents()
            self.assertIn(grid.MULTILINE_MARKER, dialog.hint.text())
        finally:
            dialog.close()
            dialog.deleteLater()


class TableExportTests(unittest.TestCase):
    """34-D：未保存表格进入同源出稿，磁盘与复杂表格不被改写。"""

    @classmethod
    def setUpClass(cls):
        cls.shared = fixtures.scratch_dir("v34-export-shared")
        # 带有效底模的真实两章项目：DOCX 与 HTML 两条链路都能验证。
        cls.project = fixtures.two_chapter_project(cls.shared / "表格工程")
        manifest = ProjectManifest.load(cls.project)
        cls.content = cls.project / manifest.relative_content_root()
        chapters = [
            path for path in sorted(cls.content.rglob("*.md")) if not path.name.startswith("_")
        ]
        cls.chapter = chapters[0]
        cls.chapter_rel = cls.chapter.relative_to(cls.content).as_posix()
        cls.chapter.write_text("# 1.1 目的\n\n原始正文。\n", encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        _cleanup(cls.shared)

    def test_unsaved_table_reaches_export_without_writing_disk(self):
        from doc_tool.application.intake_contract import (
            FORMAT_HTML, SOURCE_MODE_CURRENT_BUFFER, ExportRequest,
        )
        from doc_tool.application.project_export import run_project_export

        work = fixtures.scratch_dir("v34-export")
        try:
            chapter = self.chapter
            before = chapter.read_bytes()
            buffer_text = "### 1.1 目的\n\n" + SAMPLE_TABLE + "\n"
            report = run_project_export(
                ExportRequest(
                    project_root=str(self.project), formats=[FORMAT_HTML],
                    source_mode=SOURCE_MODE_CURRENT_BUFFER, destination=str(work / "out"),
                ),
                buffer_texts={self.chapter_rel: buffer_text},
                skip_word_refresh=True,
            )
            html = Path(report.result_for(FORMAT_HTML).path).read_text(encoding="utf-8")
            self.assertIn("=SUM(A1:A2)", html, "未保存表格必须进入同源出稿")
            self.assertIn("0012", html)
            self.assertEqual(chapter.read_bytes(), before, "出稿不得写回源文件")
        finally:
            _cleanup(work)

    def test_unsaved_table_reaches_docx_without_writing_disk(self):
        from doc_tool.application.intake_contract import (
            FORMAT_DOCX, SOURCE_MODE_CURRENT_BUFFER, ExportRequest,
        )
        from doc_tool.application.project_export import run_project_export

        work = fixtures.scratch_dir("v34-docx")
        try:
            before = self.chapter.read_bytes()
            # 章节文件内的标题必须深于文件层级（构建规则的既有约束）。
            buffer_text = "### 1.1 目的\n\n" + SAMPLE_TABLE + "\n"
            report = run_project_export(
                ExportRequest(
                    project_root=str(self.project), formats=[FORMAT_DOCX],
                    source_mode=SOURCE_MODE_CURRENT_BUFFER, destination=str(work / "out"),
                ),
                buffer_texts={self.chapter_rel: buffer_text},
                skip_word_refresh=True,
            )
            result = report.result_for(FORMAT_DOCX)
            self.assertTrue(result is not None and result.path, "当前内容应产出可读 DOCX")
            from docx import Document

            document = Document(result.path)
            texts = [
                cell.text for table_obj in document.tables
                for row in table_obj.rows for cell in row.cells
            ]
            joined = "\n".join(texts)
            self.assertIn("0012", joined, "未保存表格必须进入 DOCX")
            self.assertIn("=SUM(A1:A2)", joined, "公式文本按字符串保留")
            self.assertEqual(self.chapter.read_bytes(), before, "出稿不得写回源文件")
        finally:
            _cleanup(work)

    def test_complex_preserved_table_is_not_rewritten(self):
        complex_text = (
            "| 合并表头 |\n"
            "| 值 |\n"
        )
        parsed = grid.parse_table(complex_text)
        self.assertFalse(parsed.ok)
        parsed_two = grid.parse_table("| a | b |\n| 1 | 2 |\n| 3 | 4 |\n")
        self.assertFalse(parsed_two.ok, "缺少分隔行时不得按普通表格改写")

    def test_large_input_measurement_records_no_truncation(self):
        rows = 1000
        columns = 20
        text = "\n".join(
            "\t".join("r{0}c{1}".format(r, c) for c in range(columns)) for r in range(rows)
        )
        started = time.monotonic()
        result = grid.parse_clipboard_text(text)
        parsed_seconds = time.monotonic() - started
        self.assertTrue(result.ok, result.error)
        self.assertEqual(result.row_count, rows)
        self.assertEqual(result.column_count, columns)
        started = time.monotonic()
        serialized = grid.model_from_matrix(result.rows).serialize()
        serialize_seconds = time.monotonic() - started
        self.assertIn("r999c19", serialized, "完整输入不得截断")
        self.assertEqual(serialized.count("\n"), rows, "表头分隔行 + 999 数据行 + 表头")
        _write_evidence("large-table-measurement.json", {
            "rows": rows, "columns": columns,
            "parseSeconds": round(parsed_seconds, 3),
            "serializeSeconds": round(serialize_seconds, 3),
            "platform": "Qt offscreen / bundled Python",
            "truncated": False,
        })
        self.assertLess(parsed_seconds, 10.0, "解析不应长时间阻塞（1,000×20 文本样例）")
        self.assertLess(serialize_seconds, 10.0, "序列化不应长时间阻塞")

    def test_grid_dialog_handles_large_table_without_losing_rows(self):
        from doc_tool.ui.content.table_grid_dialog import TableGridDialog

        app = None
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        rows = [["r{0}c{1}".format(r, c) for c in range(20)] for r in range(301)]
        model = grid.model_from_matrix(rows, has_header=True)
        dialog = TableGridDialog(model)
        try:
            self.assertEqual(dialog.table.rowCount(), 301)
            self.assertEqual(dialog.table.columnCount(), 20)
            text = dialog.markdown_text()
            self.assertIn("r299c19", text)
        finally:
            dialog.close()
            dialog.deleteLater()
            app.processEvents()


if __name__ == "__main__":
    unittest.main(verbosity=2)
