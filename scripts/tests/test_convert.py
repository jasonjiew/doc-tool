# -*- coding: utf-8 -*-
"""文档互转测试（Word/PDF/Markdown/HTML/TXT/表格/RTF/ODT）。

无 Word 也能跑：方向注册表一致性、表驱动判定矩阵（含转出格式回落）、输出
命名、覆盖与同批冲突、错误码映射、批量与取消、PDF→Word 结构核验、
Markdown→Word 渲染器、Word→Markdown 离线导出、表格/文本离线后端、
组合方向路由、页范围解析与下传、适配器入参门禁。

真机部分仅在检测到可用 Word 时执行：DOCX→PDF→DOCX 往返、Markdown→Word
实转、HTML→PDF 页范围、RTF 导入。
断言刻意不解析 PDF 字节：亿赛通 DocGuard 环境下 Word 写出的 PDF 密文落盘，
Python 侧读到的是密文，只有 Word 自己能读——因此产物只验「存在且非零」。
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, REPO_ROOT)

from doc_tool.adapters import word_convert  # noqa: E402
from doc_tool.application.convert import (  # noqa: E402
    DIRECTIONS,
    DIRECTION_BY_KIND,
    KIND_CSV_TO_PDF,
    KIND_CSV_TO_XLSX,
    KIND_DOCX_TO_HTML,
    KIND_DOCX_TO_MARKDOWN,
    KIND_DOCX_TO_PDF,
    KIND_HTML_TO_DOCX,
    KIND_HTML_TO_PDF,
    KIND_MARKDOWN_TO_DOCX,
    KIND_MARKDOWN_TO_HTML,
    KIND_ODT_TO_DOCX,
    KIND_ODT_TO_PDF,
    KIND_PDF_TO_DOCX,
    KIND_RTF_TO_DOCX,
    KIND_RTF_TO_MARKDOWN,
    KIND_RTF_TO_PDF,
    KIND_RTF_TO_TXT,
    KIND_TXT_TO_PDF,
    KIND_XLSX_TO_CSV,
    KIND_XLSX_TO_HTML,
    KIND_XLSX_TO_PDF,
    SUPPORTED_SUFFIXES,
    TARGET_MARKDOWN,
    WORD_MODE_BY_KIND,
    assets_dir_for,
    build_plan,
    convert_paths,
    default_timeout_for,
    detect_kind,
    error_for_reason,
    expand_sources,
    heading_level_counts,
    parse_page_range,
    structure_hint,
    target_for,
)
from doc_tool.domain.cancellation import CancellationToken  # noqa: E402
from doc_tool.domain.errors import (  # noqa: E402
    ConversionTargetExistsError,
    PageRangeError,
    UnsupportedConversionError,
)

TEMPLATE = Path(REPO_ROOT) / "templates" / "requirement-template.docx"


def _touch(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"stub")
    return path


def _outcome(ok=True, reason="ok", detail="", elapsed=1.5, **extra):
    class _Outcome:
        pass

    result = _Outcome()
    result.ok = ok
    result.reason = reason
    result.detail = detail
    result.elapsed_seconds = elapsed
    result.error_code = extra.get("error_code")
    result.images = extra.get("images", 0)
    result.complex_tables = extra.get("complex_tables", 0)
    return result


def _converter(**kwargs):
    """按 kind 返回预置结果的假转换器，同时记录调用参数。

    签名与真 ``word_convert.convert_document`` 一致：末参 ``with_toc`` 由批量层
    原样传下来。
    """

    calls = []

    def _convert(source, target, kind, timeout_seconds, with_toc=False):
        calls.append((Path(source), Path(target), kind, timeout_seconds, with_toc))
        return kwargs.get(kind, _outcome())

    _convert.calls = calls
    return _convert


class PlanTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_kind_by_suffix(self):
        self.assertEqual(detect_kind(self.root / "a.docx"), KIND_DOCX_TO_PDF)
        self.assertEqual(detect_kind(self.root / "a.DOC"), KIND_DOCX_TO_PDF)
        self.assertEqual(detect_kind(self.root / "a.pdf"), KIND_PDF_TO_DOCX)

    def test_unsupported_suffix(self):
        with self.assertRaises(UnsupportedConversionError) as ctx:
            detect_kind(self.root / "a.zip")
        self.assertEqual(ctx.exception.code, "E6001")

    def test_target_defaults_next_to_source(self):
        source = _touch(self.root / "docs" / "手册.docx")
        self.assertEqual(target_for(source, KIND_DOCX_TO_PDF), self.root / "docs" / "手册.pdf")
        plan = build_plan(source)
        self.assertEqual(plan.kind, KIND_DOCX_TO_PDF)
        self.assertEqual(plan.label, "Word → PDF")

    def test_target_dir_override(self):
        source = _touch(self.root / "in.pdf")
        out = self.root / "out"
        self.assertEqual(target_for(source, KIND_PDF_TO_DOCX, out), out / "in.docx")

    def test_missing_source_rejected(self):
        with self.assertRaises(UnsupportedConversionError) as ctx:
            build_plan(self.root / "nope.docx")
        self.assertEqual(ctx.exception.code, "E6001")

    def test_existing_target_needs_overwrite(self):
        source = _touch(self.root / "a.docx")
        _touch(self.root / "a.pdf")
        with self.assertRaises(ConversionTargetExistsError):
            build_plan(source)
        self.assertEqual(build_plan(source, overwrite=True).target, self.root / "a.pdf")

    def test_expand_sources_from_folder(self):
        _touch(self.root / "mix" / "one.docx")
        _touch(self.root / "mix" / "two.pdf")
        _touch(self.root / "mix" / "three.log")
        names = sorted(path.name for path in expand_sources([self.root / "mix"]))
        self.assertEqual(names, ["one.docx", "two.pdf"])


def _rewrite_docx_body(dest: Path, body_xml: str) -> None:
    """复制空白模板并把 document.xml 的 body 开头替换为 ``body_xml``。"""
    shutil.copy(TEMPLATE, dest)
    with zipfile.ZipFile(TEMPLATE) as source_package:
        entries = [
            (item.filename, source_package.read(item.filename))
            for item in source_package.infolist()
        ]
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as target_package:
        for name, data in entries:
            if name == "word/document.xml":
                text = data.decode("utf-8")
                marker_end = text.index(">", text.index("<w:body")) + 1
                data = (text[:marker_end] + body_xml + text[marker_end:]).encode("utf-8")
            target_package.writestr(name, data)


def _heading_style_id(level: int) -> str:
    from doc_tool.adapters.importer import _parse_heading_styles

    with zipfile.ZipFile(TEMPLATE) as package:
        heading_map = _parse_heading_styles(package.read("word/styles.xml"))
    return next(sid for sid, lv in heading_map.items() if lv == level)


class MarkdownDirectionTests(unittest.TestCase):
    """Word 源的「转出格式」判定与 Markdown/HTML 方向的计划构造。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_docx_target_format_decides_direction(self):
        # .docx 扩展名推不出方向：缺省 PDF，显式给 md 才转 Markdown。
        source = _touch(self.root / "手册.docx")
        self.assertEqual(detect_kind(source), KIND_DOCX_TO_PDF)
        self.assertEqual(detect_kind(source, "md"), KIND_DOCX_TO_MARKDOWN)
        self.assertEqual(detect_kind(source, TARGET_MARKDOWN), KIND_DOCX_TO_MARKDOWN)
        plan = build_plan(source, target_format="md")
        self.assertEqual(plan.target, self.root / "手册.md")
        self.assertEqual(plan.label, "Word → Markdown")

    def test_legacy_doc_rejects_markdown_target(self):
        source = _touch(self.root / "old.doc")
        with self.assertRaises(UnsupportedConversionError) as ctx:
            detect_kind(source, "md")
        self.assertEqual(ctx.exception.code, "E6001")

    def test_markdown_and_html_suffixes(self):
        self.assertEqual(detect_kind(self.root / "a.md"), KIND_MARKDOWN_TO_DOCX)
        self.assertEqual(detect_kind(self.root / "a.markdown"), KIND_MARKDOWN_TO_DOCX)
        self.assertEqual(detect_kind(self.root / "a.htm"), KIND_HTML_TO_DOCX)
        self.assertEqual(
            target_for(_touch(self.root / "a.md"), KIND_MARKDOWN_TO_DOCX),
            self.root / "a.docx",
        )

    def test_expand_sources_picks_markdown_and_html(self):
        _touch(self.root / "mix" / "a.md")
        _touch(self.root / "mix" / "b.html")
        names = sorted(p.name for p in expand_sources([self.root / "mix"]))
        self.assertEqual(names, ["a.md", "b.html"])

    def test_assets_dir_collision_needs_overwrite(self):
        source = _touch(self.root / "手册.docx")
        assets_dir_for(self.root / "手册.md").mkdir(parents=True)
        with self.assertRaises(ConversionTargetExistsError):
            build_plan(source, target_format="md")
        plan = build_plan(source, target_format="md", overwrite=True)
        self.assertEqual(plan.kind, KIND_DOCX_TO_MARKDOWN)


class MarkdownBatchTests(unittest.TestCase):
    """Markdown 方向在批量层的路由与单文件失败兜底。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_docx_to_markdown_routes_to_offline_writer(self):
        source = _touch(self.root / "手册.docx")
        seen = {}

        def _writer(src, dst):
            seen["call"] = (Path(src), Path(dst))
            return _outcome(images=2, complex_tables=1, detail="2 章")

        result = convert_paths([source], target_format="md", markdown_writer=_writer)
        self.assertTrue(result.success, result.summary())
        self.assertEqual(seen["call"][0], source)
        self.assertEqual(seen["call"][1], self.root / "手册.md")
        record = result.records[0]
        self.assertEqual(record.kind, KIND_DOCX_TO_MARKDOWN)
        # 资源落点说明要进 note，供界面与 CLI 回显。
        self.assertIn("图片 2 张", record.note)
        self.assertIn("手册_assets", record.note)

    def test_markdown_to_docx_uses_word_html_mode(self):
        source = self.root / "说明.md"
        source.write_text("# 标题一\n\n正文。\n", encoding="utf-8")
        captured = {}
        converter = _converter()

        def _spy(source_arg, target_arg, kind, timeout, with_toc=False):
            captured["html"] = Path(source_arg).read_text(encoding="utf-8")
            return converter(source_arg, target_arg, kind, timeout, with_toc)

        result = convert_paths([source], converter=_spy, with_toc=True)
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_MARKDOWN_TO_DOCX)
        call = converter.calls[0]
        self.assertEqual(call[2], "html_to_docx")
        self.assertTrue(call[4])  # with_toc 原样透传给 Word 适配层
        self.assertIn("<h1>标题一</h1>", captured["html"])

    def test_undecodable_markdown_fails_single_file_as_e6007(self):
        # 未知编码：明确报 E6007，不静默 replace 出乱码，也不折成 E6003 兜底。
        source = self.root / "bad.md"
        source.write_bytes(b"\xff\xff\xff\xff")
        converter = _converter()
        result = convert_paths([source], converter=converter)
        self.assertFalse(result.success)
        self.assertEqual(result.records[0].error_code, "E6007")
        self.assertEqual(converter.calls, [])

    def test_gbk_markdown_decoded_for_html_export(self):
        # 回归：GBK 的 .md 曾被 errors="replace" 静默转成乱码产物。
        source = self.root / "gbk.md"
        source.write_bytes("# 中文标题\n\n正文。".encode("gbk"))
        result = convert_paths([source], target_format="html", overwrite=True)
        self.assertTrue(result.success, result.summary())
        html = (self.root / "gbk.html").read_text(encoding="utf-8")
        self.assertIn("中文标题", html)

    def test_backend_crash_fails_single_file_not_batch(self):
        source = _touch(self.root / "good.md")

        def _boom(source_arg, target_arg, kind, timeout, with_toc=False):
            raise RuntimeError("COM 炸了")

        result = convert_paths([source], converter=_boom)
        self.assertEqual(result.total, 1)
        self.assertEqual(result.failed, 1)
        self.assertEqual(result.records[0].error_code, "E6003")


class DocxToMarkdownTests(unittest.TestCase):
    """Word → Markdown：纯 Python 离线链路，不需要 Word。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_converts_to_single_file_with_assets(self):
        from doc_tool.application.docx_markdown import docx_to_markdown

        h1 = _heading_style_id(1)
        body = (
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr><w:r><w:t>第一章 总体要求</w:t></w:r></w:p>'
            "<w:p><w:r><w:t>正文段落一。</w:t></w:r></w:p>"
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr><w:r><w:t>第二章 交付内容</w:t></w:r></w:p>'
            "<w:p><w:r><w:t>正文段落二。</w:t></w:r></w:p>"
        ).format(h1)
        source = self.root / "手册.docx"
        _rewrite_docx_body(source, body)

        target = self.root / "out" / "手册.md"
        outcome = docx_to_markdown(source, target)
        self.assertTrue(outcome.ok, outcome.detail)
        self.assertEqual(outcome.chapters, 2)
        self.assertTrue(target.is_file())
        text = target.read_text(encoding="utf-8")
        self.assertIn("# 第一章 总体要求", text)
        self.assertIn("# 第二章 交付内容", text)
        # 构建注记与空段占位一律剥离，读者拿到的应是干净 Markdown。
        self.assertNotIn("<!--", text)
        self.assertNotIn("EMPTY_PAR", text)
        # 模板的复杂表格无法用 Markdown 表达：原样导出为 assets 资源并在正文留链接。
        assets = assets_dir_for(target)
        self.assertTrue(assets.is_dir(), "模板含复杂表格时必须落 assets 目录")
        self.assertGreater(outcome.complex_tables, 0)
        self.assertTrue(any((assets / "tables").iterdir()))
        self.assertIn(assets.name + "/tables/", text)
        # 模板的图片都在首个 Heading 1 之前，不在正文提取范围。
        self.assertEqual(outcome.images, 0)

    def test_image_link_rewrite_strips_size_dialect(self):
        from doc_tool.application.docx_markdown import _rewrite_line

        # `=96x96` 是本项目方言，通用 Markdown 渲染器会把它当 URL 的一部分，必须去掉。
        self.assertEqual(
            _rewrite_line("![logo](images/image1.png =96x96)", "手册_assets"),
            ["![logo](手册_assets/images/image1.png)"],
        )
        self.assertEqual(
            _rewrite_line("![logo](images/image1.png)", "a"),
            ["![logo](a/images/image1.png)"],
        )

    def test_missing_headings_maps_to_e1002(self):
        from doc_tool.application.docx_markdown import docx_to_markdown

        dest = self.root / "plain.docx"
        _rewrite_docx_body(dest, "<w:p><w:r><w:t>没有标题的正文</w:t></w:r></w:p>")
        outcome = docx_to_markdown(dest, self.root / "plain.md")
        self.assertFalse(outcome.ok)
        self.assertEqual(outcome.reason, "no_headings")
        self.assertEqual(outcome.error_code, "E1002")

    def test_broken_docx_maps_to_e6003(self):
        from doc_tool.application.docx_markdown import docx_to_markdown

        outcome = docx_to_markdown(
            _touch(self.root / "junk.docx"), self.root / "junk.md"
        )
        self.assertFalse(outcome.ok)
        self.assertEqual(outcome.error_code, "E6003")


class MarkdownWordRendererTests(unittest.TestCase):
    """Markdown → Word 的 HTML 渲染：结构与排版关键点。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _render(self, md: str) -> str:
        from doc_tool.application.markdown_word import markdown_to_word_html

        return markdown_to_word_html(md, self.root, title="t")

    def test_headings_paragraphs_and_inline(self):
        html = self._render("# 大标题\n\n正文 **加粗** *斜体* ~~删除~~ `code` 文本")
        self.assertIn("<h1>大标题</h1>", html)
        self.assertIn("<b>加粗</b>", html)
        self.assertIn("<i>斜体</i>", html)
        self.assertIn("<s>删除</s>", html)
        self.assertIn("<code>code</code>", html)

    def test_fenced_code_is_single_shaded_block(self):
        html = self._render("```python\nline1\nline2\n```")
        self.assertEqual(html.count("<pre>"), 1)
        # 实测裸换行会被 Word 拆成多个段落：行间必须以 <br/> 连成一块底纹。
        self.assertIn("line1<br/>line2", html)
        self.assertIn('class="language-python"', html)

    def test_nested_lists(self):
        html = self._render("- a\n  - a1\n  - a2\n- b\n")
        self.assertIn(
            "<ul><li>a<ul><li>a1</li><li>a2</li></ul></li><li>b</li></ul>", html
        )

    def test_ordered_list(self):
        html = self._render("1. one\n2. two\n")
        self.assertIn("<ol><li>one</li><li>two</li></ol>", html)

    def test_table_header_and_alignment(self):
        html = self._render("| 左 | 右 |\n| :--- | ---: |\n| a | b |")
        self.assertIn("<th>左</th>", html)
        self.assertIn('<td style="text-align:right">b</td>', html)

    def test_missing_image_keeps_visible_alt(self):
        html = self._render("![截图](不存在/图.png)")
        self.assertIn("[截图]", html)
        self.assertNotIn("<img", html)

    def test_underscore_in_link_and_image_urls_preserved(self):
        # 回归：URL 里的下划线曾被强调正则切成斜体（assets 链接/图片路径损坏）。
        _touch(self.root / "my_pic.png")
        html = self._render(
            "[资源](手册_assets/tables/tbl_0001.xml) ![图](my_pic.png)"
        )
        self.assertIn('href="手册_assets/tables/tbl_0001.xml"', html)
        self.assertIn('src=', html)
        self.assertIn("my_pic.png", html)
        self.assertNotIn("<i>assets", html)

    def test_image_width_clamped_to_page(self):
        from doc_tool.application.markdown_word import PAGE_TEXT_WIDTH_MM, _MM_TO_PX

        _touch(self.root / "big.png")
        max_px = int(PAGE_TEXT_WIDTH_MM * _MM_TO_PX)
        html = self._render("![大图](big.png =2000x1000)")
        self.assertIn('width="{0}"'.format(max_px), html)

    def test_image_only_line_centers_and_captions(self):
        _touch(self.root / "pic.png")
        html = self._render("![说明文字](pic.png)")
        self.assertIn("text-align:center", html)
        self.assertIn('<p class="caption">说明文字</p>', html)

    def test_blockquote_and_hr(self):
        html = self._render("> 引用一句\n\n---\n")
        self.assertIn("<blockquote><p>引用一句</p></blockquote>", html)
        self.assertIn("<hr/>", html)

    def test_build_notes_dropped(self):
        html = self._render("<!-- P:sz=21;sb=240 -->\n<EMPTY_PAR/>\n正文")
        self.assertNotIn("P:sz", html)
        self.assertNotIn("EMPTY_PAR", html)
        self.assertIn("<p>正文</p>", html)

    def test_read_markdown_text_strips_bom(self):
        from doc_tool.application.markdown_word import read_markdown_text

        path = self.root / "bom.md"
        path.write_text("\ufeff# 标题\n", encoding="utf-8")
        self.assertEqual(read_markdown_text(path), "# 标题\n")


class BatchTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_mixed_batch_keeps_going(self):
        docx = _touch(self.root / "good.docx")
        other = _touch(self.root / "other.docx")
        broken = _touch(self.root / "broken.docx")

        def _selective(source, target, kind, timeout, with_toc=False):
            name = Path(source).name
            if name == "good.docx":
                return _outcome(detail="PDF 已生成")
            if name == "broken.docx":
                return _outcome(ok=False, reason=word_convert.REASON_WORD_UNAVAILABLE)
            return _outcome(ok=False, reason=word_convert.REASON_CONVERT_FAILED, detail="打不开")

        result = convert_paths([docx, other, broken], converter=_selective)
        self.assertEqual(result.total, 3)
        self.assertEqual(result.succeeded, 1)
        self.assertEqual(result.failed, 2)
        self.assertFalse(result.success)
        codes = {record.source.name: record.error_code for record in result.records}
        self.assertEqual(codes["good.docx"], None)
        self.assertEqual(codes["other.docx"], "E6003")
        self.assertEqual(codes["broken.docx"], "E3001")
        # 汇总码取输入序里第一个失败；记录顺序始终等于输入顺序。
        self.assertEqual(result.error_code, "E6003")
        self.assertEqual(
            [record.source.name for record in result.records],
            ["good.docx", "other.docx", "broken.docx"],
        )

    def test_planning_failure_does_not_block_batch(self):
        docx = _touch(self.root / "ok.docx")
        weird = _touch(self.root / "note.log")
        result = convert_paths([weird, docx], converter=_converter())
        self.assertEqual(result.succeeded, 1)
        self.assertEqual(result.failed, 1)
        self.assertEqual(
            [record.source.name for record in result.records],
            ["note.log", "ok.docx"],
        )
        failed = [record for record in result.records if not record.ok][0]
        self.assertEqual(failed.error_code, "E6001")
        self.assertIn("涉及错误码：E6001", result.summary())

    def test_same_output_within_batch_is_rejected(self):
        # a.doc 与 a.docx 都落到 a.pdf：第二个必须显式失败而不是静默覆盖。
        docx = _touch(self.root / "a.docx")
        doc = _touch(self.root / "a.doc")
        result = convert_paths([docx, doc], converter=_converter())
        self.assertEqual(result.succeeded, 1)
        self.assertEqual(result.failed, 1)
        self.assertEqual(result.records[-1].error_code, "E6002")

    def test_progress_and_timeout_passed_through(self):
        docx = _touch(self.root / "b.docx")
        seen = []
        converter = _converter()
        convert_paths(
            [docx],
            converter=converter,
            timeout_seconds=42.0,
            on_progress=lambda done, total, record: seen.append((done, total)),
        )
        self.assertEqual(seen, [(1, 1)])
        self.assertEqual(converter.calls[0][3], 42.0)

    def test_cancel_stops_before_next_file(self):
        first = _touch(self.root / "c1.docx")
        second = _touch(self.root / "c2.docx")
        token = CancellationToken()
        converter = _converter()

        def _cancel_after_first(source, target, kind, timeout, with_toc=False):
            outcome = _converter()(source, target, kind, timeout)
            token.request_cancel()
            return outcome

        result = convert_paths(
            [first, second], converter=_cancel_after_first, cancel_token=token
        )
        self.assertTrue(result.cancelled)
        self.assertEqual(result.total, 1)
        self.assertIn("已取消", result.summary())

    def test_reason_to_error_code(self):
        self.assertEqual(
            error_for_reason(word_convert.REASON_TIMEOUT).code, "E6004"
        )
        self.assertEqual(
            error_for_reason(word_convert.REASON_OUTPUT_MISSING).code, "E6005"
        )
        self.assertEqual(
            error_for_reason(word_convert.REASON_WORD_UNAVAILABLE).code, "E3001"
        )
        self.assertEqual(error_for_reason("whatever").code, "E6003")

    def test_timeout_defaults_per_direction(self):
        docx = _touch(self.root / "keep.docx")
        pdf = _touch(self.root / "scan.pdf")
        converter = _converter()
        convert_paths([docx, pdf], converter=converter)
        by_kind = {kind: timeout for _, _, kind, timeout, _toc in converter.calls}
        self.assertEqual(
            by_kind[KIND_DOCX_TO_PDF], default_timeout_for(KIND_DOCX_TO_PDF)
        )
        self.assertEqual(
            by_kind[KIND_PDF_TO_DOCX], default_timeout_for(KIND_PDF_TO_DOCX)
        )
        self.assertGreater(
            default_timeout_for(KIND_PDF_TO_DOCX),
            default_timeout_for(KIND_DOCX_TO_PDF),
        )


class StructureHintTests(unittest.TestCase):
    """PDF→Word 结构核验：合成一份带 Heading 段落的 docx 来驱动真实解析路径。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _rewrite_document_body(self, dest: Path, paragraph_xml: str) -> None:
        _rewrite_docx_body(dest, paragraph_xml)

    def test_counts_heading_paragraphs_by_level(self):
        from doc_tool.adapters.importer import _parse_heading_styles

        with zipfile.ZipFile(TEMPLATE) as package:
            heading_map = _parse_heading_styles(package.read("word/styles.xml"))
        level1 = next(sid for sid, level in heading_map.items() if level == 1)
        level2 = next(sid for sid, level in heading_map.items() if level == 2)
        body = (
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr></w:p>'
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr></w:p>'
            '<w:p><w:pPr><w:pStyle w:val="{1}"/></w:pPr></w:p>'
        ).format(level1, level2)
        dest = self.root / "with_headings.docx"
        self._rewrite_document_body(dest, body)

        counts = heading_level_counts(dest)
        self.assertEqual(counts.get(1), 2)
        self.assertEqual(counts.get(2), 1)
        hint = structure_hint(dest)
        self.assertIn("已识别标题", hint)
        self.assertIn("1 级 2 个", hint)

    def test_blank_template_warns_about_missing_styles(self):
        # 空白模板只有样式定义、没有 Heading 段落：正是 PDF 还原后的典型形态。
        self.assertEqual(heading_level_counts(TEMPLATE), {})
        self.assertIn("未核验到标题样式", structure_hint(TEMPLATE))

    def test_unreadable_file_returns_empty_counts(self):
        broken = _touch(self.root / "junk.docx")
        self.assertEqual(heading_level_counts(broken), {})


class AdapterGuardTests(unittest.TestCase):
    """不启动 Word 的入参门禁与分类逻辑。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_unknown_mode_rejected(self):
        with self.assertRaises(ValueError):
            word_convert.convert_document(self.root / "a", self.root / "b", "pdf_to_xps")

    def test_missing_source_reported(self):
        outcome = word_convert.convert_document(
            self.root / "missing.docx", self.root / "out.pdf", KIND_DOCX_TO_PDF
        )
        self.assertFalse(outcome.ok)
        self.assertEqual(outcome.reason, word_convert.REASON_CONVERT_FAILED)

    def test_empty_output_reported(self):
        empty = self.root / "empty.pdf"
        empty.write_bytes(b"")
        self.assertIn("为空", word_convert._missing_output(empty))
        missing = self.root / "nothing.pdf"
        self.assertIn("未生成", word_convert._missing_output(missing))
        self.assertEqual(word_convert._missing_output(_touch(self.root / "ok.pdf")), "")

    def test_failure_classification(self):
        self.assertEqual(
            word_convert._classify_failure(Exception("no module named win32com")),
            word_convert.REASON_WORD_UNAVAILABLE,
        )
        self.assertEqual(
            word_convert._classify_failure(Exception("Word 拒绝访问")),
            word_convert.REASON_WORD_UNAVAILABLE,
        )
        self.assertEqual(
            word_convert._classify_failure(Exception("转换炸了")),
            word_convert.REASON_CONVERT_FAILED,
        )


class DirectionRegistryTests(unittest.TestCase):
    """方向注册表：一致性约束与表驱动的判定矩阵。

    新增方向只改注册表与这里的期望——矩阵逐行断言能立刻暴露「注册了但判定
    不出来」或「判定出来但没注册」的漂移。
    """

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_registry_consistency(self):
        seen_pairs = set()
        for row in DIRECTIONS:
            self.assertTrue(row.label, row.kind)
            self.assertTrue(row.target_suffix.startswith("."), row.kind)
            self.assertTrue(row.suffixes, row.kind)
            if row.backend == "word":
                self.assertIn(row.word_mode, word_convert.CONVERT_MODES)
            pair = (row.suffixes, row.target_format)
            self.assertNotIn(pair, seen_pairs, row.kind)
            seen_pairs.add(pair)
            defaults = [
                item for item in DIRECTIONS
                if item.suffixes == row.suffixes and item.is_default
            ]
            self.assertEqual(len(defaults), 1, "每个源族必须有且只有一个缺省方向")

    def test_matrix_detects_every_direction(self):
        for row in DIRECTIONS:
            source = _touch(self.root / ("matrix" + row.suffixes[0]))
            self.assertEqual(
                detect_kind(source, row.target_format), row.kind, row.kind
            )

    def test_family_default_without_explicit_target(self):
        for suffix in SUPPORTED_SUFFIXES:
            source = _touch(self.root / ("d" + suffix))
            rows = [row for row in DIRECTIONS if suffix in row.suffixes]
            default = next(row for row in rows if row.is_default)
            self.assertEqual(detect_kind(source), default.kind, suffix)

    def test_inapplicable_target_falls_back_to_family_default(self):
        # 「转换为」下拉对没有该方向的源无害：.pdf 源选 PDF 仍转 Word，
        # .md 源选 CSV 仍按缺省转 Word（矩阵外的组合不报错）。
        self.assertEqual(
            detect_kind(_touch(self.root / "s.pdf"), "pdf"), KIND_PDF_TO_DOCX
        )
        self.assertEqual(
            detect_kind(_touch(self.root / "s.md"), "csv"), KIND_MARKDOWN_TO_DOCX
        )

    def test_legacy_doc_targets_beyond_pdf_rejected(self):
        for target_format in ("md", "html", "docx"):
            with self.assertRaises(UnsupportedConversionError) as ctx:
                detect_kind(_touch(self.root / "old.doc"), target_format)
            self.assertEqual(ctx.exception.code, "E6001")


class TimeoutMappingTests(unittest.TestCase):
    """超时回归：kind 必须先经注册表翻译成 mode 再查适配器超时表。"""

    def test_word_kinds_resolve_through_mode(self):
        # 回归：曾因两套键名错位，markdown_to_docx 静默用了 docx_to_pdf 的兜底超时。
        for kind, mode in WORD_MODE_BY_KIND.items():
            self.assertEqual(
                default_timeout_for(kind),
                word_convert.default_timeout_seconds(mode),
                kind,
            )
        self.assertEqual(
            default_timeout_for(KIND_MARKDOWN_TO_DOCX),
            word_convert.default_timeout_seconds("html_to_docx"),
        )

    def test_offline_directions_use_registry_budget(self):
        direction = DIRECTION_BY_KIND[KIND_MARKDOWN_TO_HTML]
        self.assertEqual(default_timeout_for(KIND_MARKDOWN_TO_HTML), direction.offline_timeout)
        self.assertLess(
            default_timeout_for(KIND_MARKDOWN_TO_HTML),
            default_timeout_for(KIND_PDF_TO_DOCX),
        )

    def test_composite_word_directions_timeout_like_inner_word_mode(self):
        # 回归：组合方向内部真正调 Word，超时必须按内部 Word 模式取值；
        # 曾全按 offline_timeout（120s）取值，Word 还在排版就被判超时。
        for kind, mode in (
            (KIND_MARKDOWN_TO_DOCX, "html_to_docx"),
            (KIND_XLSX_TO_PDF, "html_to_pdf"),
            (KIND_CSV_TO_PDF, "html_to_pdf"),
            (KIND_TXT_TO_PDF, "html_to_pdf"),
            (KIND_RTF_TO_MARKDOWN, "import_to_docx"),
            (KIND_RTF_TO_TXT, "import_to_docx"),
        ):
            self.assertEqual(
                default_timeout_for(kind),
                word_convert.default_timeout_seconds(mode),
                kind,
            )
        offline_budget = DIRECTION_BY_KIND[KIND_MARKDOWN_TO_HTML].offline_timeout
        for kind in (KIND_XLSX_TO_PDF, KIND_RTF_TO_MARKDOWN):
            self.assertGreater(default_timeout_for(kind), offline_budget)


class PageRangeTests(unittest.TestCase):
    """页范围解析与整批下传。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_parse_valid_ranges(self):
        self.assertIsNone(parse_page_range(None))
        self.assertIsNone(parse_page_range(""))
        self.assertIsNone(parse_page_range("  "))
        self.assertEqual(parse_page_range("3"), (3, 3))
        self.assertEqual(parse_page_range("1-5"), (1, 5))
        self.assertEqual(parse_page_range(" 2 - 4 "), (2, 4))

    def test_parse_invalid_ranges(self):
        for bad in ("5-2", "0-3", "abc", "1;5", "1--2", "-3"):
            with self.assertRaises(PageRangeError) as ctx:
                parse_page_range(bad)
            self.assertEqual(ctx.exception.code, "E6008")

    def test_page_range_passed_to_word_converter(self):
        source = _touch(self.root / "p.docx")

        def _spy(source_arg, target_arg, kind, timeout, with_toc=False, page_range=None):
            seen.append(page_range)
            return _outcome()

        seen = []
        convert_paths([source], converter=_spy, page_range="2-3")
        self.assertEqual(seen, [(2, 3)])

    def test_no_page_range_keeps_five_param_signature(self):
        source = _touch(self.root / "q.docx")
        converter = _converter()
        convert_paths([source], converter=converter)
        self.assertEqual(len(converter.calls), 1)

    def test_invalid_page_range_fails_whole_batch(self):
        source = _touch(self.root / "r.docx")
        converter = _converter()
        result = convert_paths([source], converter=converter, page_range="9-x")
        self.assertEqual(result.failed, 1)
        self.assertEqual(result.records[0].error_code, "E6008")
        self.assertEqual(converter.calls, [])


class TableBackendTests(unittest.TestCase):
    """表格离线后端（XLSX/CSV）：读取降级、互转与 PDF 组合路由。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _write_xlsx(self, path: Path):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "报价"
        sheet.append(["物料", "数量", "单价"])
        sheet.append(["螺钉", 3.0, 1.5])
        sheet.append([None, None, None])
        sheet["A4"] = "合并备注"
        sheet.merge_cells("A4:B4")
        second = workbook.create_sheet("附表")
        second.append(["备注"])
        second.append(["第二张表"])
        workbook.save(str(path))

    def test_xlsx_to_csv_default_direction(self):
        source = self.root / "表.xlsx"
        self._write_xlsx(source)
        result = convert_paths([source])
        self.assertTrue(result.success, result.summary())
        record = result.records[0]
        self.assertEqual(record.kind, KIND_XLSX_TO_CSV)
        target = self.root / "表.csv"
        self.assertTrue(target.is_file())
        text = target.read_text(encoding="utf-8-sig")
        self.assertIn("物料,数量,单价", text)
        self.assertIn("螺钉,3,1.5", text)
        # 多余工作表不静默丢弃：产物说明里要点名。
        self.assertIn("其余 1 个工作表", record.detail)

    def test_merged_cells_filled_with_top_left_value(self):
        source = self.root / "表.xlsx"
        self._write_xlsx(source)
        convert_paths([source], overwrite=True)
        text = (self.root / "表.csv").read_text(encoding="utf-8-sig")
        self.assertIn("合并备注,合并备注", text)

    def test_corrupted_xlsx_maps_to_e6006(self):
        source = _touch(self.root / "junk.xlsx")
        result = convert_paths([source])
        self.assertFalse(result.success)
        self.assertEqual(result.records[0].error_code, "E6006")

    def test_undecodable_csv_maps_to_e6007(self):
        # CSV 编码解不出是 E6007，不能折进「Word 转换失败」（E6003）的兜底。
        source = self.root / "bad.csv"
        source.write_bytes(b"\xff\xff\xff\xff")
        result = convert_paths([source], overwrite=True)
        self.assertFalse(result.success)
        self.assertEqual(result.records[0].error_code, "E6007")
        self.assertIn("编码", result.records[0].detail)

    def test_csv_to_xlsx_roundtrip(self):
        source = self.root / "表.csv"
        source.write_text("物料,数量\n螺钉,3\n", encoding="utf-8-sig")
        result = convert_paths([source])
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_CSV_TO_XLSX)
        from openpyxl import load_workbook

        workbook = load_workbook(str(self.root / "表.xlsx"))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("物料", "数量"))
        self.assertEqual(rows[1], ("螺钉", "3"))

    def test_csv_gbk_decoded(self):
        source = self.root / "gbk.csv"
        source.write_bytes("物料,数量\n螺钉,3\n".encode("gbk"))
        result = convert_paths([source], overwrite=True)
        self.assertTrue(result.success, result.summary())
        from openpyxl import load_workbook

        workbook = load_workbook(str(self.root / "gbk.xlsx"))
        self.assertEqual(workbook.active["A2"].value, "螺钉")

    def test_xlsx_to_html_lists_every_sheet(self):
        source = self.root / "表.xlsx"
        self._write_xlsx(source)
        result = convert_paths([source], target_format="html")
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_XLSX_TO_HTML)
        text = (self.root / "表.html").read_text(encoding="utf-8")
        self.assertIn("<h1>报价</h1>", text)
        self.assertIn("<h1>附表</h1>", text)
        self.assertIn("<th>物料</th>", text)

    def test_xlsx_to_pdf_routes_through_word_html_mode(self):
        source = self.root / "表.xlsx"
        self._write_xlsx(source)
        captured = {}

        def _spy(source_arg, target_arg, kind, timeout, with_toc=False):
            captured["mode"] = kind
            captured["html"] = Path(source_arg).read_text(encoding="utf-8")
            return _outcome(detail="PDF 已生成")

        result = convert_paths([source], target_format="pdf", converter=_spy)
        self.assertTrue(result.success, result.summary())
        self.assertEqual(captured["mode"], "html_to_pdf")
        self.assertIn("<h1>报价</h1>", captured["html"])
        self.assertIn("工作表", result.records[0].detail)

    def test_csv_to_pdf_routes_through_word_html_mode(self):
        source = self.root / "表.csv"
        source.write_text("物料,数量\n", encoding="utf-8-sig")
        converter = _converter(**{"html_to_pdf": _outcome(detail="PDF 已生成")})
        result = convert_paths(
            [source], target_format="pdf", converter=converter, overwrite=True
        )
        self.assertTrue(result.success, result.summary())
        self.assertEqual(converter.calls[0][2], "html_to_pdf")


class TextDirectionTests(unittest.TestCase):
    """TXT → PDF 组合方向与 DOCX → 纯文本。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_txt_to_pdf_routes_through_word_html_mode(self):
        source = self.root / "说明.txt"
        source.write_text("第一行\n\n第二行\n", encoding="utf-8")
        captured = {}

        def _spy(source_arg, target_arg, kind, timeout, with_toc=False):
            captured["mode"] = kind
            captured["html"] = Path(source_arg).read_text(encoding="utf-8")
            return _outcome(detail="PDF 已生成")

        result = convert_paths([source], converter=_spy)
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_TXT_TO_PDF)
        self.assertEqual(captured["mode"], "html_to_pdf")
        self.assertIn("第一行", captured["html"])
        self.assertIn("文本 3 行", result.records[0].detail)

    def test_txt_gbk_decoded(self):
        source = self.root / "gbk.txt"
        source.write_bytes("中文内容\n".encode("gbk"))
        captured = {}

        def _spy(source_arg, target_arg, kind, timeout, with_toc=False):
            captured["html"] = Path(source_arg).read_text(encoding="utf-8")
            return _outcome()

        result = convert_paths([source], converter=_spy)
        self.assertTrue(result.success, result.summary())
        self.assertIn("中文内容", captured["html"])

    def test_undecodable_txt_maps_to_e6007(self):
        source = self.root / "bad.txt"
        source.write_bytes(b"\xff\xff\xff\xff")
        converter = _converter()
        result = convert_paths([source], converter=converter)
        self.assertFalse(result.success)
        self.assertEqual(result.records[0].error_code, "E6007")
        self.assertEqual(converter.calls, [])

    def test_docx_to_text_extracts_paragraphs_and_tables(self):
        from docx import Document

        source = self.root / "样例.docx"
        document = Document()
        document.add_paragraph("第一章 总体要求")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "列A"
        table.cell(0, 1).text = "列B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
        document.save(str(source))

        from doc_tool.application.text_convert import docx_to_text

        target = self.root / "out" / "样例.txt"
        outcome = docx_to_text(source, target)
        self.assertTrue(outcome.ok, outcome.detail)
        text = target.read_text(encoding="utf-8")
        self.assertIn("第一章 总体要求", text)
        self.assertIn("列A | 列B", text)
        self.assertIn("1 | 2", text)


class MarkdownHtmlExportTests(unittest.TestCase):
    """Markdown/Word → 单文件 HTML（离线）。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_markdown_to_html_reuses_renderer(self):
        source = self.root / "说明.md"
        source.write_text("# 标题一\n\n正文 **加粗**。\n", encoding="utf-8")
        result = convert_paths([source], target_format="html")
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_MARKDOWN_TO_HTML)
        text = (self.root / "说明.html").read_text(encoding="utf-8")
        self.assertIn("<h1>标题一</h1>", text)
        self.assertIn("<b>加粗</b>", text)

    def test_docx_to_html_with_assets_and_no_leftover_md(self):
        h1 = _heading_style_id(1)
        body = (
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr><w:r><w:t>第一章 总体要求</w:t></w:r></w:p>'
            "<w:p><w:r><w:t>正文段落一。</w:t></w:r></w:p>"
        ).format(h1)
        source = self.root / "手册.docx"
        _rewrite_docx_body(source, body)
        result = convert_paths([source], output_dir=self.root / "out", target_format="html")
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_DOCX_TO_HTML)
        target = self.root / "out" / "手册.html"
        text = target.read_text(encoding="utf-8")
        self.assertIn("第一章 总体要求", text)
        # 模板的复杂表格走 assets 机制，HTML 里留提示链接；不残留 Markdown 中间文件。
        self.assertTrue(assets_dir_for(target).is_dir())
        self.assertIn("_assets/", text)
        self.assertFalse((self.root / "out" / "手册.md").exists())
        self.assertIn("_assets", result.records[0].note)


class ImportDirectionTests(unittest.TestCase):
    """RTF/ODT 方向：路由与 RTF 组合方向（导入 → 离线导出）。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_rtf_and_odt_routing(self):
        rtf = _touch(self.root / "样例.rtf")
        odt = _touch(self.root / "样例.odt")
        converter = _converter()
        convert_paths([rtf], converter=converter)
        self.assertEqual(converter.calls[0][2], "import_to_pdf")
        convert_paths([rtf], target_format="docx", converter=converter, overwrite=True)
        self.assertEqual(converter.calls[1][2], "import_to_docx")
        convert_paths([odt], target_format="docx", converter=converter, overwrite=True)
        self.assertEqual(converter.calls[2][2], "import_to_docx")

    def _word_fake_writing_docx(self, calls):
        """假 Word 转换器：把带 Heading 1 的真 DOCX 写到目标（临时 docx 路径）。"""
        h1 = _heading_style_id(1)
        body = (
            '<w:p><w:pPr><w:pStyle w:val="{0}"/></w:pPr><w:r><w:t>第一章 总体要求</w:t></w:r></w:p>'
            "<w:p><w:r><w:t>正文段落一。</w:t></w:r></w:p>"
        ).format(h1)

        def _convert(source_arg, target_arg, kind, timeout, with_toc=False):
            calls.append(kind)
            _rewrite_docx_body(Path(target_arg), body)
            return _outcome(detail="导入 Word 完成，2 段")

        return _convert

    def test_rtf_to_markdown_composite(self):
        source = _touch(self.root / "样例.rtf")
        calls = []
        result = convert_paths(
            [source], output_dir=self.root / "out",
            target_format="md", converter=self._word_fake_writing_docx(calls),
        )
        self.assertTrue(result.success, result.summary())
        self.assertEqual(calls, ["import_to_docx"])
        self.assertEqual(result.records[0].kind, KIND_RTF_TO_MARKDOWN)
        md = self.root / "out" / "样例.md"
        text = md.read_text(encoding="utf-8")
        self.assertIn("# 第一章 总体要求", text)
        self.assertIn("_assets", result.records[0].note)

    def test_rtf_to_text_composite(self):
        source = _touch(self.root / "样例.rtf")
        calls = []
        result = convert_paths(
            [source], output_dir=self.root / "out",
            target_format="txt", converter=self._word_fake_writing_docx(calls),
        )
        self.assertTrue(result.success, result.summary())
        self.assertEqual(result.records[0].kind, KIND_RTF_TO_TXT)
        text = (self.root / "out" / "样例.txt").read_text(encoding="utf-8")
        self.assertIn("第一章 总体要求", text)
        self.assertIn("导入 Word 完成", result.records[0].detail)

    def test_rtf_composite_word_failure_propagates(self):
        source = _touch(self.root / "样例.rtf")

        def _fail(source_arg, target_arg, kind, timeout, with_toc=False):
            return _outcome(ok=False, reason=word_convert.REASON_TIMEOUT)

        result = convert_paths([source], target_format="md", converter=_fail)
        self.assertFalse(result.success)
        self.assertEqual(result.records[0].error_code, "E6004")


def _word_static() -> bool:
    """不 spawn 子进程的静态门禁。

    这里绝不能用 ``dispatch_check=True``：该检查经 ``multiprocessing`` spawn
    子进程，子进程会重新导入本模块，若在导入期探测就会递归 spawn，并把门禁
    误判成「无 Word」。实启动探测放到 ``setUp``（只在父进程执行）。
    """
    try:
        from doc_tool.application.word_check import check_word_available

        return check_word_available(dispatch_check=False).available
    except Exception:
        return False


@unittest.skipUnless(_word_static(), "缺少 pywin32 或非交互式桌面会话，跳过真机互转")
class RealWordConversionTests(unittest.TestCase):
    """真机往返：DOCX→PDF→DOCX，走的就是用户那一条代码路径。"""

    def setUp(self):
        from doc_tool.application.word_check import check_word_available

        if not check_word_available(dispatch_check=True).available:
            self.skipTest("本机 Word 未能通过 COM 启动探测")
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.docx = self.root / "手册.docx"
        shutil.copy(TEMPLATE, self.docx)

    def tearDown(self):
        self._tmp.cleanup()

    def test_roundtrip(self):
        to_pdf = convert_paths([self.docx])
        self.assertTrue(to_pdf.success, to_pdf.summary() + " " + str(to_pdf.records))
        pdf = to_pdf.records[0].target
        self.assertTrue(pdf.is_file())
        # DocGuard 下 Python 读 Word 写出的 PDF 只会拿到密文，因此不解析字节。
        self.assertGreater(pdf.stat().st_size, 1024)
        self.assertGreater(to_pdf.records[0].elapsed_seconds, 0.0)

        # 同目录回转会和原始 手册.docx 同名：绝不覆盖用户原稿，必须先被拦下。
        collide = convert_paths([pdf])
        self.assertFalse(collide.success)
        self.assertEqual(collide.records[0].error_code, "E6002")

        back = convert_paths([pdf], self.root / "back")
        self.assertTrue(back.success, back.summary() + " " + str(back.records))
        rebuilt = back.records[0].target
        self.assertEqual(rebuilt.suffix, ".docx")
        self.assertGreater(rebuilt.stat().st_size, 1024)
        # PDF 还原必须给出结构核验说明，且文案与实际解析出的层级数量一致
        # （实测 21 页模板只回来 7 个 1 级标题、2 级及以下为 0：层级不保证，
        # 但界面不能拿一套话术、核验结果另说一套）。
        note = back.records[0].note
        counts = heading_level_counts(rebuilt)
        if counts:
            self.assertIn("已识别标题", note)
            for level, number in counts.items():
                self.assertIn("{0} 级 {1} 个".format(level, number), note)
        else:
            self.assertIn("未核验到标题样式", note)

    def test_second_run_needs_overwrite(self):
        convert_paths([self.docx])
        again = convert_paths([self.docx])
        self.assertFalse(again.success)
        self.assertEqual(again.records[0].error_code, "E6002")
        forced = convert_paths([self.docx], overwrite=True)
        self.assertTrue(forced.success)

    def test_markdown_to_docx_and_back(self):
        source = self.root / "样品.md"
        source.write_text(
            "# 接口说明\n\n正文段落，含 **加粗** 与 `code`。\n\n"
            "- 列表一\n- 列表二\n\n"
            "| 列A | 列B |\n| :--- | ---: |\n| 1 | 2 |\n\n"
            "```python\nprint('hi')\nprint('bye')\n```\n",
            encoding="utf-8",
        )
        # 去程：Markdown → Word（真机 Word 完成 HTML 导入、图片内嵌与排版收口）。
        result = convert_paths([source], output_dir=self.root / "word")
        self.assertTrue(result.success, result.summary() + " " + str(result.records))
        docx = result.records[0].target
        self.assertGreater(docx.stat().st_size, 1024)
        counts = heading_level_counts(docx)
        self.assertTrue(counts.get(1), "Markdown 一级标题未映射成 Heading 1")
        self.assertIn("产物标题", result.records[0].note)

        # 回程：Word 产物 → Markdown（离线）。.docx 明文落盘，Python 侧可读。
        back = convert_paths(
            [docx], output_dir=self.root / "back", target_format="md"
        )
        self.assertTrue(back.success, back.summary() + " " + str(back.records))
        md = back.records[0].target
        self.assertTrue(md.is_file())
        self.assertIn("# 接口说明", md.read_text(encoding="utf-8"))

    def test_html_to_pdf_page_range(self):
        source = self.root / "三页.html"
        source.write_text(
            "<html><head><meta charset='utf-8'></head><body>"
            "<h1>第一页</h1><p>内容一</p>"
            "<p style='page-break-before:always'>第二页</p>"
            "<p style='page-break-before:always'>第三页</p>"
            "</body></html>",
            encoding="utf-8",
        )
        result = convert_paths(
            [source], output_dir=self.root / "out",
            target_format="pdf", page_range="2-2",
        )
        self.assertTrue(result.success, result.summary() + " " + str(result.records))
        record = result.records[0]
        self.assertEqual(record.kind, KIND_HTML_TO_PDF)
        pdf = record.target
        # DocGuard 下只验存在与体量，不解析字节。
        self.assertGreater(pdf.stat().st_size, 1024)
        self.assertIn("（第 2-2 页）", record.detail)

    def test_rtf_imports(self):
        rtf = self.root / "样品.rtf"
        rtf.write_text(
            "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 SimSun;}}"
            "\\f0\\fs24 RTF Probe Heading\\par RTF Probe Body\\par}",
            encoding="ascii",
        )
        to_docx = convert_paths(
            [rtf], output_dir=self.root / "out", target_format="docx"
        )
        self.assertTrue(to_docx.success, to_docx.summary() + " " + str(to_docx.records))
        docx = to_docx.records[0].target
        self.assertEqual(docx.suffix, ".docx")
        self.assertGreater(docx.stat().st_size, 1024)
        # 导入方向同样给结构核验说明（模板无标题时应给出未核验提示）。
        self.assertTrue(to_docx.records[0].note)

        to_txt = convert_paths(
            [rtf], output_dir=self.root / "out2", target_format="txt"
        )
        self.assertTrue(to_txt.success, to_txt.summary() + " " + str(to_txt.records))
        text = to_txt.records[0].target.read_text(encoding="utf-8")
        self.assertIn("RTF Probe Heading", text)
        self.assertIn("RTF Probe Body", text)


if __name__ == "__main__":
    unittest.main()
