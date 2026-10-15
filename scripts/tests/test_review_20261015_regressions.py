# -*- coding: utf-8 -*-
"""2026-10-15 review regressions for lossless import, merge and document output."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
for candidate in (str(ROOT), str(ROOT / "scripts")):
    if candidate not in sys.path:
        sys.path.insert(0, candidate)

from core_fixtures import two_chapter_project
from doc_tool.application.content.reimport import ChapterChange, ReimportService, merge_table_block_three_way
from doc_tool.application.content.writer import ContentWriter
from doc_tool.application.import_project import ImportRequest, import_first_time
from doc_tool.adapters.fidelity import SEVERITY_BLOCK
from doc_tool.adapters.roundtrip import RoundtripIssue, RoundtripReport
from doc_tool.domain.manifest import ProjectManifest
from doc_tool.domain.paths import ProjectPaths
from doc_tool.domain.split_rules import PROFILE_AUTO_RULES
from doc_tool.domain.blocks import KIND_COMPLEX_TABLE, parse_blocks
from doc_tool.domain.captions import build_registry, resolve_inline_references
from doc_tool.domain.merge_policy import (
    SCOPED_MERGE_POLICY,
    matches_chapter_scope,
    shift_inline_headings_for_depth,
)
from doc_tool.domain.smart_features import (
    apply_complex_table_proxy_edits,
    bind_imported_captions_and_references,
    compute_smart_column_widths,
    format_complex_table_proxy_lines,
)
from docx_common import _is_event_text_equivalent, iter_chapter_entries
from build_docx import reference_segments, set_active_registry


class ReviewContentSafetyTests(unittest.TestCase):
    HEADER = "| key | value |\n| --- | --- |\n"

    def test_table_merge_uses_row_identity_after_reordering(self):
        base = self.HEADER + "| A | old |\n| B | old |"
        local = self.HEADER + "| B | old |\n| A | old |"
        incoming = self.HEADER + "| A | new |\n| B | old |"
        merged, clean = merge_table_block_three_way(base, local, incoming)
        self.assertTrue(clean)
        self.assertIn("| B | old |\n| A | new |", merged)

    def test_table_merge_rejects_conflicting_caption_and_alignment(self):
        base = self.HEADER + "| A | old |"
        for local, incoming in (
            ("Table: local\n" + base, "Table: incoming\n" + base),
            (base.replace("--- | ---", ":--- | ---"),
             base.replace("--- | ---", "---: | ---")),
        ):
            with self.subTest(local=local):
                baseline = "Table: base\n" + base if local.startswith("Table:") else base
                _, clean = merge_table_block_three_way(baseline, local, incoming)
                self.assertFalse(clean)

    def test_table_merge_keeps_unilateral_alignment_change(self):
        base = self.HEADER + "| A | old |"
        local = base.replace("| A | old |", "| A | local |")
        incoming = base.replace("--- | ---", "---: | ---")
        merged, clean = merge_table_block_three_way(base, local, incoming)
        self.assertTrue(clean)
        self.assertIn("| ---: | --- |", merged)
        self.assertIn("| A | local |", merged)

    def test_table_merge_does_not_truncate_extra_cells(self):
        base = self.HEADER + "| A | old | must survive |"
        local = base.replace("| A | old |", "| A | local |")
        _, clean = merge_table_block_three_way(base, local, base)
        self.assertFalse(clean)

    @staticmethod
    def _table():
        return ET.fromstring(
            '<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr>'
            '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>first</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>second</w:t></w:r></w:p>'
            '</w:tc></w:tr></w:tbl>'
        )

    @staticmethod
    def _parsed_proxy(table):
        text = "<!-- TABLE:1:tbl_0001.xml -->\n" + "\n".join(format_complex_table_proxy_lines(table))
        return next(b.proxy_rows for b in parse_blocks(text).blocks if b.kind == KIND_COMPLEX_TABLE)

    def test_unedited_multiline_proxy_preserves_original_xml(self):
        table = self._table()
        before = ET.tostring(table)
        proxy = self._parsed_proxy(table)
        self.assertFalse(apply_complex_table_proxy_edits(table, proxy))
        self.assertEqual(ET.tostring(table), before)

    def test_proxy_roundtrip_preserves_backslashes_and_pipes(self):
        table = self._table()
        text_node = next(table.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        text_node.text = r"C:\\folder\name\|literal"
        before = ET.tostring(table)
        self.assertFalse(apply_complex_table_proxy_edits(table, self._parsed_proxy(table)))
        self.assertEqual(ET.tostring(table), before)

    def test_proxy_rejects_shape_and_merged_placeholder_edits_before_writing(self):
        for rows in (
            [["changed", ""], ["extra", ""]],
            [["changed"]],
            [["changed", "cannot edit a merged continuation"]],
        ):
            with self.subTest(rows=rows):
                table = self._table()
                before = ET.tostring(table)
                with self.assertRaises(ValueError):
                    apply_complex_table_proxy_edits(table, rows)
                self.assertEqual(ET.tostring(table), before)

    def test_caption_binding_preserves_fenced_examples(self):
        for fence in ("```", "~~~~"):
            with self.subTest(fence=fence):
                example = [fence + "text", "见图 1", "![fake](fake.png)", "图 2 示例代码", fence]
                lines = ["![real](real.png)", "图 1 正式题注", "", *example, "", "正文见图 1。"]
                bind_imported_captions_and_references(lines)
                self.assertEqual(lines[3:3 + len(example)], example)
                self.assertIn("@fig-1", lines[-1])

    def test_caption_binding_handles_multiline_table_entries_without_losing_indices(self):
        lines = [
            "| key | value |\n| --- | --- |\n| A | B |", "",
            "![real](real.png)", "图 1 正式题注", "", "见图 1。",
        ]
        bind_imported_captions_and_references(lines)
        self.assertIn("Figure:", lines[3])
        self.assertIn("@fig-1", lines[5])

    def test_heading_shift_keeps_info_string_inside_code_fence(self):
        markdown = "```text\n```python\n# code must remain\n```\n# prose\n"
        shifted, delta = shift_inline_headings_for_depth(markdown, 2, enabled=True)
        self.assertEqual(delta, 2)
        self.assertEqual(shifted, "```text\n```python\n# code must remain\n```\n### prose\n")

    def test_wide_tables_have_positive_widths_with_exact_total(self):
        for count in (30, 40, 80):
            with self.subTest(count=count):
                widths = compute_smart_column_widths([["x"] * count], 9000)
                self.assertEqual(len(widths), count)
                self.assertEqual(sum(widths), 9000)
                self.assertTrue(all(width > 0 for width in widths), widths)

    def test_numeric_section_scope_retains_ancestors_and_descendants_only(self):
        self.assertTrue(matches_chapter_scope("第2章 设计", ["2.1"]))
        self.assertTrue(matches_chapter_scope("第2章 设计/2.1 接口/2.1.1 输入.md", ["2.1"]))
        self.assertFalse(matches_chapter_scope("第2章 设计/2.2 其他.md", ["2.1"]))
        self.assertFalse(matches_chapter_scope("第1章 引言", ["2.1"]))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "第2章 设计").mkdir()
            for name in ("2.1 接口.md", "2.2 其他.md"):
                (root / "第2章 设计" / name).write_text("Body.\n", encoding="utf-8")
            entries = list(iter_chapter_entries({
                "paths": {"content_root": str(root)},
                "mergePolicy": SCOPED_MERGE_POLICY.with_scope(["2.1"]),
            }))
            self.assertEqual([Path(e.path).name for e, _ in entries], ["第2章 设计", "2.1 接口.md"])

    def test_roundtrip_reference_normalization_keeps_business_text(self):
        self.assertFalse(_is_event_text_equivalent("参见图 1 保留全部内容。", "参见图 2 删除全部内容。"))
        self.assertTrue(_is_event_text_equivalent("参见图 1-1 所示。", "参见图 2 所示。"))

    def test_ascii_reference_does_not_consume_following_chinese_prose(self):
        registry = build_registry([parse_blocks("![diagram](image.png)\nFigure: Diagram {#fig-one}\n")])
        expected = "如" + registry.entries["fig-one"].display + "所示。"
        self.assertEqual(resolve_inline_references("如@fig-one所示。", registry), expected)
        set_active_registry(registry)
        try:
            segments = reference_segments("如@fig-one所示。")
            self.assertEqual([kind for kind, _ in segments], ["text", "ref", "text"])
            self.assertEqual(segments[-1], ("text", "所示。"))
        finally:
            set_active_registry(None)


    def test_ambiguous_duplicate_row_keys_are_not_silently_merged(self):
        base = self.HEADER + "| A | first |\n| A | second |"
        local = base.replace("first", "local")
        incoming = base.replace("second", "incoming")
        _, clean = merge_table_block_three_way(base, local, incoming)
        self.assertFalse(clean)

    def test_reimport_resource_rewrite_preserves_fenced_literals(self):
        text = "```markdown\n![literal](images/img_0001.png)\n<!-- TABLE:1:tbl_0001.xml -->\n```\n"
        mapping = {"img_0001.png": "images/img_0002.png", "tbl_0001.xml": "tbl_0002.xml"}
        self.assertEqual(ReimportService._rewrite_incoming_refs(text, mapping), text)

    def test_generated_code_fence_is_longer_than_literal_delimiters(self):
        from doc_tool.adapters.importer import _CODE_LINE_TAG, _cluster_smart_code_lines
        original = ["```python", "print('literal markdown')", "```"]
        output = _cluster_smart_code_lines([_CODE_LINE_TAG + line for line in original])
        parsed = parse_blocks("\n".join(output))
        self.assertEqual(len(parsed.blocks), 1)
        self.assertEqual(len(parsed.code_blocks), 1)
        self.assertEqual(parsed.code_blocks[0].text, "\n".join(original))

    def test_clean_title_policy_preserves_explicit_mode(self):
        from doc_tool.domain.split_rules import NUMBERING_CLEAN_TITLE, SplitPolicy
        policy = SplitPolicy(numbering_mode=NUMBERING_CLEAN_TITLE)
        self.assertEqual(policy.numbering_mode, NUMBERING_CLEAN_TITLE)
        self.assertEqual(SplitPolicy.from_dict(policy.to_dict()).numbering_mode, NUMBERING_CLEAN_TITLE)

    def test_heading_promotion_only_warns_when_title_text_is_preserved(self):
        from doc_tool.adapters.roundtrip import _classify_error
        from doc_tool.adapters.fidelity import SEVERITY_WARN
        from validate_docx import Event
        source = [Event("P", "第1章 Introduction", "body[0]")]
        rebuilt = [Event("H", (1, "Introduction"), "body[0]")]
        error = "#0 类型不一致: P -> H"
        self.assertEqual(_classify_error(error, source, rebuilt).severity, SEVERITY_BLOCK)
        self.assertEqual(_classify_error(error, source, rebuilt, heading_promotion_mode="preserve_semantic").severity, SEVERITY_WARN)
        changed = [Event("H", (1, "Deleted important words"), "body[0]")]
        self.assertEqual(_classify_error(error, source, changed, heading_promotion_mode="preserve_semantic").severity, SEVERITY_BLOCK)


class ReimportContinuityTests(unittest.TestCase):
    """Real service transactions, including repeated imports and fault injection."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="review-reimport-")
        self.addCleanup(self.temporary.cleanup)
        self.work = Path(self.temporary.name)
        self.project = two_chapter_project(self.work / "project")
        self.manifest = ProjectManifest.load(self.project)
        self.paths = ProjectPaths(self.project)
        self.content = self.paths.resolve(self.manifest.relative_content_root())
        self.rel = "第1章 引言/1.1 目的.md"
        self.base = "First stable paragraph.\n\nLast stable paragraph.\n"
        (self.content / self.rel).write_text(self.base, encoding="utf-8")
        self.incoming = self.work / "incoming-content"
        shutil.copytree(self.content, self.incoming)
        self.source = self.work / "incoming-source.docx"
        shutil.copy2(self.paths.source_docx, self.source)
        self.service = ReimportService(self.manifest, self.paths)
        self.incoming_image = None
        self.asset_name = "img_0001.png"
        self._save_base()

    def _save_base(self):
        paths = [path for path in self.content.rglob("*.md") if path.is_file()]
        self.base_file = self.paths.state_dir / "reimport_base.json"
        self.base_file.write_text(json.dumps({
            "files": {path.relative_to(self.content).as_posix(): hashlib.sha1(path.read_bytes()).hexdigest() for path in paths},
            "contents": {path.relative_to(self.content).as_posix(): path.read_text(encoding="utf-8") for path in paths},
        }, ensure_ascii=False), encoding="utf-8")
        self.original_base = self.base_file.read_bytes()

    def _extract(self, _source, destination):
        shutil.copytree(self.incoming, destination, dirs_exist_ok=True)
        if self.incoming_image is not None:
            images = destination.parent / "images"
            images.mkdir(parents=True, exist_ok=True)
            (images / self.asset_name).write_bytes(self.incoming_image)

    @staticmethod
    def _png(color):
        stream = BytesIO()
        Image.new("RGB", (2, 2), color).save(stream, "PNG")
        return stream.getvalue()

    def test_repeated_unchanged_source_keeps_local_edits(self):
        local = self.base.replace("First stable", "First locally edited")
        incoming = self.base.replace("Last stable", "Last externally edited")
        (self.content / self.rel).write_text(local, encoding="utf-8")
        (self.incoming / self.rel).write_text(incoming, encoding="utf-8")
        for _ in range(2):
            result = self.service.reimport(self.source, extractor=self._extract)
            self.assertTrue(result.success, result.message)
            actual = (self.content / self.rel).read_text(encoding="utf-8")
            self.assertIn("First locally edited", actual)
            self.assertIn("Last externally edited", actual)
        baseline = json.loads(self.base_file.read_text(encoding="utf-8"))
        self.assertEqual(baseline["contents"][self.rel], incoming)

    def test_merged_local_image_is_not_remapped_as_an_incoming_image(self):
        images = self.paths.images_dir(self.manifest.documentType)
        images.mkdir(parents=True, exist_ok=True)
        original_image = self._png("red")
        (images / self.asset_name).write_bytes(original_image)
        self.incoming_image = self._png("blue")
        local = self.base.replace("\n\n", "\n\n![local](images/" + self.asset_name + ")\n\n", 1)
        incoming = self.base + "\n![incoming](images/" + self.asset_name + ")\n"
        (self.content / self.rel).write_text(local, encoding="utf-8")
        (self.incoming / self.rel).write_text(incoming, encoding="utf-8")
        for _ in range(2):
            result = self.service.reimport(self.source, extractor=self._extract)
            self.assertTrue(result.success, result.message)
            actual = (self.content / self.rel).read_text(encoding="utf-8")
            references = dict(re.findall(r"!\[(local|incoming)\]\(images/([^\s)]+)\)", actual))
            self.assertEqual(set(references), {"local", "incoming"}, actual)
            self.assertNotEqual(references["local"], references["incoming"])
            self.assertEqual((images / references["local"]).read_bytes(), original_image)
            self.assertEqual((images / references["incoming"]).read_bytes(), self.incoming_image)
        self.assertEqual(len(list(images.glob("*.png"))), 2)

    def test_asset_only_change_is_visible_in_readonly_preview_and_applied(self):
        images = self.paths.images_dir(self.manifest.documentType)
        images.mkdir(parents=True, exist_ok=True)
        (images / self.asset_name).write_bytes(self._png("red"))
        self.incoming_image = self._png("blue")
        text = self.base + "\n![diagram](images/" + self.asset_name + ")\n"
        (self.content / self.rel).write_text(text, encoding="utf-8")
        (self.incoming / self.rel).write_text(text, encoding="utf-8")
        self._save_base()
        before_images = {path.name: path.read_bytes() for path in images.iterdir() if path.is_file()}
        preview, error = self.service.preview(self.source, extractor=self._extract)
        self.addCleanup(self.service._release_preview_staging)
        self.assertIsNotNone(preview, error)
        entry = next(item for item in preview.entries if item.rel_path == self.rel)
        self.assertEqual(entry.change, "modified")
        self.assertTrue(entry.applicable)
        self.assertEqual({path.name: path.read_bytes() for path in images.iterdir() if path.is_file()}, before_images)
        self.assertEqual((self.content / self.rel).read_text(encoding="utf-8"), text)
        result = self.service.reimport(self.source, extractor=self._extract)
        self.assertTrue(result.success, result.message)
        actual = (self.content / self.rel).read_text(encoding="utf-8")
        target = re.search(r"!\[diagram\]\(images/([^\s)]+)\)", actual).group(1)
        self.assertEqual((images / target).read_bytes(), self.incoming_image)

    def test_failed_rename_deletion_rolls_back_created_destination(self):
        destination = "第1章 引言/1.1 目的更新.md"
        (self.incoming / self.rel).rename(self.incoming / destination)
        with patch.object(ContentWriter, "delete_file", return_value=SimpleNamespace(written=False, error="synthetic delete failure")):
            result = self.service.reimport(self.source, extractor=self._extract)
        self.assertFalse(result.success)
        self.assertIn("synthetic delete failure", result.message)
        self.assertEqual((self.content / self.rel).read_text(encoding="utf-8"), self.base)
        self.assertFalse((self.content / destination).exists())
        self.assertEqual(self.base_file.read_bytes(), self.original_base)

    def test_failed_create_does_not_rollback_a_prior_user_create(self):
        rel = "第1章 引言/1.2 Local.md"
        writer = self.service._writer()
        self.assertTrue(writer.create_file(rel, "Prior local content.\n").written)
        (self.incoming / rel).write_text("Incoming replacement.\n", encoding="utf-8")
        stale_comparison = [ChapterChange(rel, "added", use_new=True)]
        with patch("doc_tool.application.content.reimport.compare_chapters", return_value=stale_comparison):
            result = self.service.reimport(self.source, extractor=self._extract)
        self.assertFalse(result.success)
        self.assertTrue((self.content / rel).is_file())
        self.assertEqual((self.content / rel).read_text(encoding="utf-8"), "Prior local content.\n")
        self.assertIn(("create", rel), self.service._writer().manifest.keys())

    def test_variant_and_explicit_scope_are_intersected(self):
        from doc_tool.application.pipeline import run_pipeline
        (self.project / "variants.yml").write_text(
            "schemaVersion: 1\nvariants:\n  - variantId: restricted\n    name: Restricted\n    include:\n      - 第2章 设计\n",
            encoding="utf-8",
        )
        result = run_pipeline(self.manifest, self.paths, skip_word_refresh=True,
                              variant_id="restricted", chapter_scope=["第1章 引言"])
        self.assertFalse(result.success)
        self.assertIsNone(result.output_path)

    def test_empty_variant_fallback_is_not_silent(self):
        from doc_tool.application.pipeline import run_pipeline
        (self.project / "variants.yml").write_text(
            "schemaVersion: 1\nvariants:\n  - variantId: missing\n    name: Missing\n    include:\n      - nonexistent-chapter\n",
            encoding="utf-8",
        )
        result = run_pipeline(self.manifest, self.paths, skip_word_refresh=True, variant_id="missing")
        self.assertTrue(result.success, [(event.stage, event.status, event.detail) for event in result.events])
        self.assertTrue(any("有效章节为空" in event.detail for event in result.events))

    def test_failed_publish_restores_original_source_baseline(self):
        local = self.base.replace("First stable", "First locally edited")
        (self.content / self.rel).write_text(local, encoding="utf-8")
        (self.incoming / self.rel).write_text(self.base.replace("Last stable", "Last external"), encoding="utf-8")
        with patch.object(self.manifest, "save", side_effect=OSError("synthetic manifest failure")):
            result = self.service.reimport(self.source, extractor=self._extract)
        self.assertFalse(result.success)
        self.assertEqual((self.content / self.rel).read_text(encoding="utf-8"), local)
        self.assertEqual(json.loads(self.base_file.read_text(encoding="utf-8"))["contents"][self.rel], self.base)
        self.assertEqual(self.base_file.read_bytes(), self.original_base)

    def test_explicit_rollback_restores_previous_source_baseline(self):
        local = self.base.replace("First stable", "First locally edited")
        (self.content / self.rel).write_text(local, encoding="utf-8")
        (self.incoming / self.rel).write_text(self.base.replace("Last stable", "Last external"), encoding="utf-8")
        result = self.service.reimport(self.source, extractor=self._extract)
        self.assertTrue(result.success, result.message)
        self.assertEqual(self.service.rollback_last(), [])
        self.assertEqual((self.content / self.rel).read_text(encoding="utf-8"), local)
        self.assertEqual(json.loads(self.base_file.read_text(encoding="utf-8"))["contents"][self.rel], self.base)
        self.assertEqual(self.base_file.read_bytes(), self.original_base)

    def test_heuristic_heading_import_does_not_bypass_missing_content_gate(self):
        loss = RoundtripReport(4, 3, (RoundtripIssue(SEVERITY_BLOCK, "synthetic missing body", "#count"),))
        target = self.work / "must-not-publish"
        request = ImportRequest(
            source_docx=self.source, target_project_root=target,
            document_type="general", document_no="REVIEW-001",
            document_name="Synthetic review fixture", document_version="1.0",
            split_policy=PROFILE_AUTO_RULES,
        )
        with patch("doc_tool.application.import_project._run_roundtrip_check", return_value=loss):
            result = import_first_time(request)
        self.assertFalse(result.success)
        self.assertEqual(result.error_code, "E2004")
        self.assertFalse(target.exists())


class RoundtripSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="review-roundtrip-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_code_container_alignment_does_not_hide_later_prose_loss(self):
        from docx import Document
        from doc_tool.adapters.roundtrip import roundtrip_diff

        source = Document()
        source.add_heading("Review", 1)
        for text in ("code first", "code second", "Important original prose", "Tail"):
            source.add_paragraph(text)
        rebuilt = Document()
        rebuilt.add_heading("Review", 1)
        rebuilt.add_table(rows=1, cols=1).cell(0, 0).text = "code first\ncode second"
        rebuilt.add_paragraph("code second")  # Matches the wrong, unaligned source index.
        rebuilt.add_paragraph("Tail")
        source.save(str(self.root / "source.docx"))
        rebuilt.save(str(self.root / "rebuilt.docx"))
        self.assertTrue(roundtrip_diff(self.root / "source.docx", self.root / "rebuilt.docx").has_block)

    def test_twenty_representation_warnings_do_not_hide_later_prose_loss(self):
        from docx import Document
        from doc_tool.adapters.roundtrip import roundtrip_diff

        for filename, label in (("source.docx", "original"), ("rebuilt.docx", "changed")):
            document = Document()
            document.add_heading("Review", 1)
            for index in range(20):
                document.add_table(rows=1, cols=1).cell(0, 0).text = "Table {0}: {1}".format(index, label)
            document.add_paragraph("Important prose: " + label)
            document.save(str(self.root / filename))
        report = roundtrip_diff(self.root / "source.docx", self.root / "rebuilt.docx")
        self.assertTrue(report.has_block)
        self.assertLessEqual(len(report.issues), 20)


class StaticTocSafetyTests(unittest.TestCase):
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def _root(self, extra=""):
        from lxml import etree
        return etree.fromstring((r'''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body><w:sdt><w:sdtContent><w:p>
            <w:r><w:t>Prefix</w:t></w:r>
            <w:r><w:fldChar w:fldCharType="begin"/></w:r>
            <w:r><w:instrText> TOC \o "1-3" \h \z \u </w:instrText></w:r>
            <w:r><w:fldChar w:fldCharType="separate"/></w:r>
            <w:r><w:t>Stale cache</w:t></w:r>CACHE
            <w:r><w:fldChar w:fldCharType="end"/></w:r>
            <w:r><w:t>Suffix</w:t></w:r>
          </w:p></w:sdtContent></w:sdt>
          <w:p><w:bookmarkStart w:id="7" w:name="one"/><w:r><w:t>Body one</w:t></w:r><w:bookmarkEnd w:id="7"/></w:p>
          <w:p><w:bookmarkStart w:id="8" w:name="two"/><w:r><w:t>Body two</w:t></w:r><w:bookmarkEnd w:id="8"/></w:p>
          </w:body></w:document>''').replace("CACHE", extra).encode())

    def test_cache_is_refreshable_balanced_and_idempotent(self):
        from lxml import etree
        from doc_tool.kernel_shared.toc_cache import update_static_toc_cache
        root = self._root()
        warnings = []
        headings = [(1, "New one", "one"), (2, "New two", "two")]
        self.assertEqual(update_static_toc_cache(root, headings, warnings.append), 2)
        first = etree.tostring(root)
        self.assertEqual(update_static_toc_cache(root, headings, warnings.append), 2)
        self.assertEqual(etree.tostring(root), first)
        self.assertEqual([node.get(self.W + "anchor") for node in root.iter(self.W + "hyperlink")], ["one", "two"])
        text = [node.text for node in root.iter(self.W + "t")]
        for expected in ("Prefix", "Suffix", "New one", "New two", "Body one", "Body two"):
            self.assertEqual(text.count(expected), 1)
        self.assertNotIn("Stale cache", text)
        self.assertEqual(text.count("?"), 2)
        balance = 0
        for field in root.iter(self.W + "fldChar"):
            kind = field.get(self.W + "fldCharType")
            balance += (kind == "begin") - (kind == "end")
            self.assertGreaterEqual(balance, 0)
        self.assertEqual(balance, 0)
        self.assertTrue(warnings)

    def test_unsupported_cache_layouts_are_unchanged_with_warning(self):
        from lxml import etree
        from doc_tool.kernel_shared.toc_cache import update_static_toc_cache
        for mode in ("bookmark", "comment", "section", "custom_switch", "invalid_range", "unrecognized_range"):
            with self.subTest(mode=mode):
                extra = '<w:bookmarkStart w:id="9" w:name="protected"/>' if mode == "bookmark" else ""
                if mode == "comment":
                    extra = '<w:r><w:commentReference w:id="2"/></w:r>'
                root = self._root(extra)
                if mode == "section":
                    paragraph = root.find(".//" + self.W + "sdtContent/" + self.W + "p")
                    properties = etree.Element(self.W + "pPr")
                    etree.SubElement(properties, self.W + "sectPr")
                    paragraph.insert(0, properties)
                elif mode == "custom_switch":
                    root.find(".//" + self.W + "instrText").text = ' TOC \\b subset '
                elif mode == "invalid_range":
                    root.find(".//" + self.W + "instrText").text = ' TOC \\o "9-2" '
                elif mode == "unrecognized_range":
                    root.find(".//" + self.W + "instrText").text = ' TOC \\o OOPS '
                before = etree.tostring(root)
                warnings = []
                self.assertEqual(update_static_toc_cache(root, [(1, "New one", "one")], warnings.append), 0)
                self.assertEqual(etree.tostring(root), before)
                self.assertTrue(warnings)

    def test_missing_toc_is_noop_and_incomplete_field_is_preserved(self):
        from lxml import etree
        from doc_tool.kernel_shared.toc_cache import update_static_toc_cache
        root = etree.Element(self.W + "document")
        etree.SubElement(root, self.W + "body")
        before = etree.tostring(root)
        warnings = []
        self.assertEqual(update_static_toc_cache(root, [], warnings.append), 0)
        self.assertEqual(etree.tostring(root), before)
        self.assertEqual(warnings, [])
        root = self._root()
        end = next(node for node in root.iter(self.W + "fldChar") if node.get(self.W + "fldCharType") == "end")
        end.getparent().remove(end)
        before = etree.tostring(root)
        self.assertEqual(update_static_toc_cache(root, [(1, "New one", "one")], warnings.append), 0)
        self.assertEqual(etree.tostring(root), before)
        self.assertTrue(warnings)


class PublicMarkdownExportSafetyTests(unittest.TestCase):
    def test_public_export_omits_project_only_complex_table_proxy_markers(self):
        from docx import Document
        from doc_tool.application.docx_markdown import docx_to_markdown

        with tempfile.TemporaryDirectory(prefix="review-public-markdown-") as temporary:
            root = Path(temporary)
            source = root / "source.docx"
            target = root / "manual.md"
            document = Document()
            document.add_heading("Export", 1)
            document.add_paragraph("Visible body.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
            table.cell(1, 0).text = "First"
            table.cell(1, 1).text = "Second"
            document.save(str(source))
            result = docx_to_markdown(source, target)
            self.assertTrue(result.ok, result.detail)
            self.assertEqual(result.complex_tables, 1)
            text = target.read_text(encoding="utf-8")
            self.assertIn("Visible body.", text)
            self.assertNotIn("<!--", text)
            self.assertNotIn("TABLE_PROXY_", text)
            self.assertIn("manual_assets/tables/tbl_0001.xml", text)
            self.assertTrue((root / "manual_assets/tables/tbl_0001.xml").is_file())


class LegacyHeadingImportSafetyTests(unittest.TestCase):
    def test_h3_first_import_preserves_legacy_source_style_and_baseline(self):
        from docx import Document
        from doc_tool.application.import_project import ImportRequest, import_first_time
        from doc_tool.adapters.kernel import build_with_project, validate_with_project
        from doc_tool.domain.manifest import ProjectManifest
        from doc_tool.domain.paths import ProjectPaths

        with tempfile.TemporaryDirectory(prefix="review-legacy-heading-") as temporary:
            root = Path(temporary)
            source = root / "source.docx"
            target = root / "project"
            document = Document()
            document.add_heading("Top-level source style", 3)
            document.add_paragraph("Visible original body.")
            document.save(str(source))
            result = import_first_time(ImportRequest(
                source_docx=str(source), target_project_root=str(target),
                document_type="general", document_no="REVIEW-LEGACY",
                document_name="Legacy heading", document_version="1.0",
                allow_missing_headings=True,
            ))
            self.assertTrue(result.success, result.error_code)
            manifest = ProjectManifest.load(target / "project.yml")
            self.assertEqual(manifest.headingStyles[1], "Heading3")
            paths = ProjectPaths(target)
            build_with_project(manifest, paths)
            self.assertTrue(validate_with_project(manifest, paths, baseline=True))


if __name__ == "__main__":
    unittest.main()
